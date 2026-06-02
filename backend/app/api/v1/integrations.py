import io
from typing import Annotated

from fastapi import APIRouter, Depends, File, HTTPException, Request, UploadFile, status
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.deps import get_current_user, get_db
from app.core.limiter import limiter
from app.core.security import decrypt_secret, encrypt_secret, verify_password
from app.core.upload_validation import validate_excel_upload
from app.models.integration import Integration
from app.models.user import User
from app.schemas.integration import EXCHANGE_PROVIDERS, IntegrationCreate, IntegrationOut
from app.services.audit import AuditAction, log_audit

router = APIRouter(prefix="/integrations", tags=["integrations"])

CurrentUser = Annotated[User, Depends(get_current_user)]
DbSession = Annotated[AsyncSession, Depends(get_db)]

_XLSX_MEDIA = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_XLSX_HEADERS = {"Content-Disposition": "attachment; filename=kripto-api-anahtarlari.xlsx"}
_EXPORT_RESOURCE = "integration:export.xlsx"


@router.get("", response_model=list[IntegrationOut])
async def list_integrations(
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[AsyncSession, Depends(get_db)],
):
    result = await db.execute(select(Integration).where(Integration.user_id == current_user.id))
    return result.scalars().all()


@router.post("", response_model=IntegrationOut, status_code=status.HTTP_201_CREATED)
async def add_integration(
    request: Request,
    payload: IntegrationCreate,
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[AsyncSession, Depends(get_db)],
):
    result = await db.execute(
        select(Integration).where(
            Integration.user_id == current_user.id,
            Integration.provider == payload.provider,
        )
    )
    existing = result.scalar_one_or_none()
    if existing:
        existing.encrypted_key = encrypt_secret(payload.api_key)
        existing.encrypted_secret = encrypt_secret(payload.api_secret) if payload.api_secret else None
        existing.is_active = True
        await log_audit(
            db,
            request,
            action=AuditAction.INTEGRATION_ADD,
            user_id=current_user.id,
            resource=f"integration:{payload.provider}",
            extra={"updated": True},
        )
        await db.commit()
        await db.refresh(existing)
        return existing

    integration = Integration(
        user_id=current_user.id,
        provider=payload.provider,
        encrypted_key=encrypt_secret(payload.api_key),
        encrypted_secret=encrypt_secret(payload.api_secret) if payload.api_secret else None,
    )
    db.add(integration)
    await log_audit(
        db,
        request,
        action=AuditAction.INTEGRATION_ADD,
        user_id=current_user.id,
        resource=f"integration:{payload.provider}",
        extra={"updated": False},
    )
    await db.commit()
    await db.refresh(integration)
    return integration


@router.delete("/{provider}", status_code=status.HTTP_204_NO_CONTENT)
async def remove_integration(
    provider: str,
    request: Request,
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[AsyncSession, Depends(get_db)],
):
    result = await db.execute(
        select(Integration).where(
            Integration.user_id == current_user.id,
            Integration.provider == provider,
        )
    )
    integration = result.scalar_one_or_none()
    if not integration:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Entegrasyon bulunamadı")
    await db.delete(integration)
    await log_audit(
        db,
        request,
        action=AuditAction.INTEGRATION_DELETE,
        user_id=current_user.id,
        resource=f"integration:{provider}",
    )
    await db.commit()


class IntegrationExportRequest(BaseModel):
    """Tam (decrypt edilmiş) API key export'u için şifre doğrulama gövdesi."""

    password: str


def _build_integrations_xlsx(rows: list[Integration]) -> io.BytesIO:
    """Entegrasyonları decrypt edilmiş API key/secret ile xlsx'e yazar."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Kripto API Anahtarları"

    warning_text = (
        "GUVENLIK UYARISI: Bu dosya borsa API anahtarlarinizi ACIK (decrypt edilmis) "
        "icerir. Sizmasi halinde hesaplariniza erisilebilir. E-postayla paylasmayin, "
        "bulut deposunda sifresiz tutmayin; kullandiktan sonra silin. API key'leri "
        "salt-okunur + IP kisitli olusturmaniz onerilir."
    )
    wcell = ws.cell(row=1, column=1, value=warning_text)
    wcell.font = Font(bold=True, color="C53030")
    wcell.fill = PatternFill("solid", fgColor="FED7D7")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=3)

    headers = ["Borsa", "API Key", "API Secret"]
    hfill = PatternFill("solid", fgColor="7C3AED")
    hfont = Font(bold=True, color="FFFFFF")
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=2, column=col, value=h)
        c.fill = hfill
        c.font = hfont
        c.alignment = Alignment(horizontal="center")

    for ridx, intg in enumerate(rows, 3):
        ws.cell(row=ridx, column=1, value=intg.provider)
        ws.cell(row=ridx, column=2, value=decrypt_secret(intg.encrypted_key) if intg.encrypted_key else "")
        ws.cell(row=ridx, column=3, value=decrypt_secret(intg.encrypted_secret) if intg.encrypted_secret else "")

    for col, width in zip("ABC", [16, 60, 60]):
        ws.column_dimensions[col].width = width

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


@router.post(
    "/export",
    responses={403: {"description": "Şifre hatalı — API anahtarları verilmez"}},
)
@limiter.limit("5/minute")
async def export_integrations_full(body: IntegrationExportRequest, request: Request, current_user: CurrentUser, db: DbSession):
    """Borsa API anahtarlarını ACIK (decrypt) Excel olarak dışa aktarır — şifre doğrulamalı.

    Şifre yanlış/eksikse 403 ve anahtarlar VERILMEZ. API key/secret çok hassastır
    (hesap erişimi); bu yüzden her zaman şifre arkasında + audit'lenir (`integration.export`),
    rate limit 5/dk. xpub export'u (wallets) ile aynı güvenlik modeli.
    """
    if not verify_password(body.password, current_user.password_hash):
        await log_audit(
            db,
            request,
            action=AuditAction.INTEGRATION_EXPORT,
            user_id=current_user.id,
            resource=_EXPORT_RESOURCE,
            extra={"result": "wrong_password"},
        )
        await db.commit()
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Şifre hatalı — API anahtarları verilmedi.")

    result = await db.execute(select(Integration).where(Integration.user_id == current_user.id, Integration.encrypted_key.is_not(None)))
    rows = result.scalars().all()
    buf = _build_integrations_xlsx(rows)
    await log_audit(
        db,
        request,
        action=AuditAction.INTEGRATION_EXPORT,
        user_id=current_user.id,
        resource=_EXPORT_RESOURCE,
        extra={"count": len(rows)},
    )
    await db.commit()
    return StreamingResponse(buf, media_type=_XLSX_MEDIA, headers=_XLSX_HEADERS)


def _parse_integration_row(row: tuple) -> dict | None:
    """Excel satırını (Borsa, API Key, API Secret) doğrular; geçersizse None."""
    provider = str(row[0]).strip().lower() if row and row[0] else ""
    if not provider or provider == "none" or provider not in EXCHANGE_PROVIDERS:
        return None
    api_key = str(row[1]).strip() if len(row) > 1 and row[1] else ""
    if not api_key or api_key == "none":
        return None
    api_secret = str(row[2]).strip() if len(row) > 2 and row[2] else None
    if api_secret == "none":
        api_secret = None
    return {"provider": provider, "api_key": api_key, "api_secret": api_secret}


@router.post("/import", response_model=list[IntegrationOut])
async def import_integrations(file: Annotated[UploadFile, File()], request: Request, current_user: CurrentUser, db: DbSession):
    """Excel'den borsa API anahtarlarını içe aktarır (provider bazında upsert).

    Dosya `_build_integrations_xlsx` formatında (Borsa / API Key / API Secret). Mevcut
    sağlayıcı güncellenir, yenisi eklenir; dosyada olmayan entegrasyonlara DOKUNULMAZ
    (replace-all değil — kimlik bilgisi kaybını önler).
    """
    from openpyxl import load_workbook

    content = await validate_excel_upload(file)
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
    except Exception as exc:
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="Dosya okunamadı") from exc

    parsed: list[dict] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        item = _parse_integration_row(row)
        if item is not None:
            parsed.append(item)
    if not parsed:
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="Geçerli API anahtarı bulunamadı")

    existing_rows = await db.execute(select(Integration).where(Integration.user_id == current_user.id))
    by_provider = {i.provider: i for i in existing_rows.scalars().all()}
    for p in parsed:
        intg = by_provider.get(p["provider"])
        if intg is None:
            intg = Integration(user_id=current_user.id, provider=p["provider"])
            db.add(intg)
            by_provider[p["provider"]] = intg
        intg.encrypted_key = encrypt_secret(p["api_key"])
        intg.encrypted_secret = encrypt_secret(p["api_secret"]) if p["api_secret"] else None
        intg.is_active = True

    await log_audit(
        db,
        request,
        action=AuditAction.INTEGRATION_ADD,
        user_id=current_user.id,
        resource=_EXPORT_RESOURCE,
        extra={"imported": len(parsed), "via": "import"},
    )
    await db.commit()
    result = await db.execute(select(Integration).where(Integration.user_id == current_user.id))
    return result.scalars().all()


@router.post("/sync", status_code=status.HTTP_202_ACCEPTED)
async def sync_integrations(current_user: Annotated[User, Depends(get_current_user)]):
    # TODO: arka planda senkronizasyon görevi başlat
    return {"detail": "Senkronizasyon başlatıldı"}
