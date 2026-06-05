import calendar
import io
import logging
from datetime import date as date_type
from datetime import datetime
from decimal import Decimal, InvalidOperation
from typing import Annotated
from zoneinfo import ZoneInfo

# Realize işlemlerinde "bugün" Türkiye saatine göre belirlenmeli — snapshot'la
# tutarlı (backend Docker UTC'de çalışır).
_ISTANBUL = ZoneInfo("Europe/Istanbul")

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile, status
from fastapi.responses import StreamingResponse
from sqlalchemy import desc, func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.deps import get_current_user, get_db
from app.core.upload_validation import validate_excel_upload
from app.models.income import Income
from app.models.recurring_income import RecurringIncome
from app.models.recurring_skip import RecurringSkip
from app.models.user import User
from app.schemas.income import (
    INCOME_CATEGORIES,
    IncomeCategoryBreakdown,
    IncomeCreate,
    IncomeDashboard,
    IncomeOut,
    IncomeSummary,
    IncomeUpdate,
    RealizeMonthRequest,
    RealizeResult,
    RecurringIncomeCreate,
    RecurringIncomeOut,
    RecurringIncomeUpdate,
)
from app.schemas.recurring import (
    RecurringPeriodsResult,
    RecurringPeriodStatus,
    RecurringUnrealizeResult,
)
from app.services import currency as currency_svc
from app.services import recurrence

logger = logging.getLogger(__name__)
router = APIRouter(prefix="/income", tags=["income"])

_NOT_FOUND_DETAIL = "Kayıt bulunamadı"
_RECURRING_NOT_FOUND = "Periyodik kayıt bulunamadı"
_OPENPYXL_MISSING = "openpyxl kütüphanesi bulunamadı"

# Türkçe label -> İngilizce key haritası (import için)
LABEL_TO_KEY: dict[str, str] = {
    "maaş": "salary",
    "serbest meslek": "freelance",
    "kira geliri": "rental",
    "temettü / faiz": "dividend",
    "temettü": "dividend",
    "ikramiye / prim": "bonus",
    "ikramiye": "bonus",
    "varlık satışı": "sale",
    "diğer": "other",
    "salary": "salary",
    "freelance": "freelance",
    "rental": "rental",
    "dividend": "dividend",
    "bonus": "bonus",
    "sale": "sale",
    "other": "other",
}


@router.get("", response_model=list[IncomeOut])
async def list_incomes(
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[AsyncSession, Depends(get_db)],
    year: Annotated[int | None, Query(ge=2020, le=2100)] = None,
    month: Annotated[int | None, Query(ge=1, le=12)] = None,
    category: Annotated[str | None, Query()] = None,
):
    stmt = select(Income).where(Income.user_id == current_user.id)
    if year is not None and month is not None:
        first_day = date_type(year, month, 1)
        last_day = date_type(year, month, calendar.monthrange(year, month)[1])
        stmt = stmt.where(Income.date >= first_day, Income.date <= last_day)
    if category:
        if category not in INCOME_CATEGORIES:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail=f"Geçersiz kategori: {category}",
            )
        stmt = stmt.where(Income.category == category)
    stmt = stmt.order_by(desc(Income.date), desc(Income.created_at))
    result = await db.execute(stmt)
    return result.scalars().all()


async def _resolve_amount_tl(amount: Decimal, currency: str) -> tuple[Decimal, Decimal | None]:
    """İşlem-anı kuruyla TL karşılığı + 1 birim kuru döner (hibrit kur, SABİT).

    TRY için exchange_rate=1 (audit basit). Diğer dövizlerde güncel kur haritası
    çekilip convert_to_tl ile amount_tl sabitlenir.
    """
    if (currency or "TRY").upper() == "TRY":
        return Decimal(amount), Decimal("1")
    rates = await currency_svc.fetch_rates()
    amount_tl = currency_svc.convert_to_tl(amount, currency, rates)
    rate = rates.get(currency.upper())
    return amount_tl, rate


@router.post("", response_model=IncomeOut, status_code=status.HTTP_201_CREATED)
async def create_income(
    payload: IncomeCreate,
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[AsyncSession, Depends(get_db)],
):
    currency = payload.currency or current_user.default_currency or "TRY"
    amount_tl, exchange_rate = await _resolve_amount_tl(payload.amount, currency)
    inc = Income(
        user_id=current_user.id,
        amount=payload.amount,
        currency=currency,
        amount_tl=amount_tl,
        exchange_rate=exchange_rate,
        category=payload.category,
        date=payload.date,
        description=payload.description.strip() if payload.description else None,
    )
    db.add(inc)
    await db.commit()
    await db.refresh(inc)
    return inc


@router.put("/{income_id}", response_model=IncomeOut)
async def update_income(
    income_id: int,
    payload: IncomeUpdate,
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[AsyncSession, Depends(get_db)],
):
    result = await db.execute(select(Income).where(Income.id == income_id, Income.user_id == current_user.id))
    inc = result.scalar_one_or_none()
    if not inc:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=_NOT_FOUND_DETAIL)

    if payload.amount is not None:
        inc.amount = payload.amount
    if payload.currency is not None:
        inc.currency = payload.currency
    if payload.category is not None:
        inc.category = payload.category
    if payload.date is not None:
        inc.date = payload.date
    if payload.description is not None:
        inc.description = payload.description.strip() or None

    # Tutar veya para birimi değiştiyse amount_tl yeniden sabitlenir (güncel kur).
    if payload.amount is not None or payload.currency is not None:
        inc.amount_tl, inc.exchange_rate = await _resolve_amount_tl(inc.amount, inc.currency)

    await db.commit()
    await db.refresh(inc)
    return inc


@router.delete("/{income_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_income(
    income_id: int,
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[AsyncSession, Depends(get_db)],
):
    result = await db.execute(select(Income).where(Income.id == income_id, Income.user_id == current_user.id))
    inc = result.scalar_one_or_none()
    if not inc:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=_NOT_FOUND_DETAIL)
    await db.delete(inc)
    await db.commit()


@router.get("/summary", response_model=IncomeSummary)
async def get_income_summary(
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[AsyncSession, Depends(get_db)],
    year: Annotated[int, Query(ge=2020, le=2100)],
    month: Annotated[int, Query(ge=1, le=12)],
):
    first_day = date_type(year, month, 1)
    last_day = date_type(year, month, calendar.monthrange(year, month)[1])

    # v0.3.0 çoklu para birimi: toplamlar amount_tl (işlem-anı kuruyla sabit) üzerinden.
    total_q = await db.execute(
        select(func.coalesce(func.sum(Income.amount_tl), 0), func.count(Income.id)).where(
            Income.user_id == current_user.id, Income.date >= first_day, Income.date <= last_day
        )
    )
    total, count = total_q.one()

    cat_q = await db.execute(
        select(Income.category, func.sum(Income.amount_tl), func.count(Income.id))
        .where(Income.user_id == current_user.id, Income.date >= first_day, Income.date <= last_day)
        .group_by(Income.category)
        .order_by(desc(func.sum(Income.amount_tl)))
    )
    breakdown = [IncomeCategoryBreakdown(category=cat, total=Decimal(amt), count=cnt) for cat, amt, cnt in cat_q.all()]

    return IncomeSummary(
        year=year,
        month=month,
        total=Decimal(total),
        count=count,
        by_category=breakdown,
    )


@router.get("/export", responses={500: {"description": _OPENPYXL_MISSING}})
async def export_incomes(
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[AsyncSession, Depends(get_db)],
    year: Annotated[int | None, Query(ge=2020, le=2100)] = None,
    month: Annotated[int | None, Query(ge=1, le=12)] = None,
):
    """Gelirleri Excel dosyası olarak indir."""
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        raise HTTPException(status_code=500, detail=_OPENPYXL_MISSING)

    stmt = select(Income).where(Income.user_id == current_user.id)
    if year is not None and month is not None:
        first_day = date_type(year, month, 1)
        last_day = date_type(year, month, calendar.monthrange(year, month)[1])
        stmt = stmt.where(Income.date >= first_day, Income.date <= last_day)
    stmt = stmt.order_by(Income.date, Income.created_at)

    result = await db.execute(stmt)
    incomes = result.scalars().all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Gelirler"

    # Başlık satırı: yeşil arka plan, beyaz bold
    header_fill = PatternFill(start_color="059669", end_color="059669", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    headers = ["Tarih", "Kategori", "Tutar (₺)", "Açıklama"]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for row_idx, inc in enumerate(incomes, start=2):
        ws.cell(row=row_idx, column=1, value=str(inc.date))
        ws.cell(row=row_idx, column=2, value=inc.category)
        ws.cell(row=row_idx, column=3, value=float(inc.amount))
        ws.cell(row=row_idx, column=4, value=inc.description or "")

    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 40

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=gelirler.xlsx"},
    )


def _parse_excel_date(date_val) -> date_type | None:
    """Excel hücresinden tarih parse — desteklenmeyen/boş ise None."""
    if isinstance(date_val, date_type):
        return date_val
    if isinstance(date_val, str):
        try:
            return date_type.fromisoformat(date_val.strip())
        except ValueError:
            return None
    try:
        from openpyxl.utils.datetime import from_excel

        return from_excel(date_val).date() if date_val is not None else None
    except Exception:
        return None


def _parse_amount(amount_val) -> Decimal | None:
    """Tutar parse — pozitif değilse veya geçersizse None."""
    try:
        amount = Decimal(str(amount_val)).quantize(Decimal("0.01"))
    except (InvalidOperation, TypeError):
        return None
    return amount if amount > 0 else None


def _income_from_row(row, user_id) -> Income | None:
    """Bir Excel satırından Income üretir; geçersiz satırda None döner."""
    n = len(row)
    date_val, cat_val, amount_val, desc_val = (row[i] if i < n else None for i in range(4))

    parsed_date = _parse_excel_date(date_val)
    if parsed_date is None:
        return None

    cat_str = str(cat_val).strip().lower() if cat_val is not None else ""
    category = LABEL_TO_KEY.get(cat_str)
    if not category:
        return None

    amount = _parse_amount(amount_val)
    if amount is None:
        return None

    description = str(desc_val).strip() if desc_val else None
    # Excel import TRY varsayar (sütunda para birimi yok) → amount_tl = amount.
    return Income(
        user_id=user_id,
        amount=amount,
        currency="TRY",
        amount_tl=amount,
        exchange_rate=Decimal("1"),
        category=category,
        date=parsed_date,
        description=description or None,
    )


@router.post(
    "/import",
    response_model=list[IncomeOut],
    status_code=status.HTTP_201_CREATED,
    responses={
        400: {"description": "Geçersiz Excel dosyası"},
        500: {"description": _OPENPYXL_MISSING},
    },
)
async def import_incomes(
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[AsyncSession, Depends(get_db)],
    file: Annotated[UploadFile, File()],
):
    """Excel dosyasından gelir içe aktar (append — mevcut kayıtlar silinmez)."""
    try:
        import openpyxl
    except ImportError:
        raise HTTPException(status_code=500, detail=_OPENPYXL_MISSING)

    # SEC-009 (FAZ H): magic-byte + boyut + extension dogrulamasi
    content = await validate_excel_upload(file)
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    except Exception:
        raise HTTPException(status_code=400, detail="Geçersiz Excel dosyası")

    ws = wb.active
    added: list[Income] = []

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or all(v is None for v in row):
            continue
        inc = _income_from_row(row, current_user.id)
        if inc is None:
            continue
        db.add(inc)
        added.append(inc)

    await db.commit()
    for inc in added:
        await db.refresh(inc)

    return added


# ---------------------------------------------------------------------------
# Periyodik gelir (recurring_incomes) CRUD + dashboard hesaplama
# ---------------------------------------------------------------------------
def _applies_in_month(ri: RecurringIncome, year: int, month: int) -> bool:
    """Bir periyodik gelirin verilen ay içinde geçerli olup olmadığı (ortak util)."""
    return recurrence.applies_in_month(ri, year, month)


@router.get("/recurring", response_model=list[RecurringIncomeOut])
async def list_recurring_incomes(
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[AsyncSession, Depends(get_db)],
):
    result = await db.execute(select(RecurringIncome).where(RecurringIncome.user_id == current_user.id).order_by(RecurringIncome.start_date.desc()))
    return result.scalars().all()


@router.post("/recurring", response_model=RecurringIncomeOut, status_code=status.HTTP_201_CREATED)
async def create_recurring_income(
    payload: RecurringIncomeCreate,
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[AsyncSession, Depends(get_db)],
):
    if payload.recurrence == "custom" and not payload.months:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="custom recurrence için 'months' alanı zorunlu (1-12 arası ay listesi)",
        )
    ri = RecurringIncome(
        user_id=current_user.id,
        title=payload.title,
        amount=payload.amount,
        currency=payload.currency or current_user.default_currency or "TRY",
        category=payload.category,
        recurrence=payload.recurrence,
        months=payload.months,
        day_of_month=payload.day_of_month,
        start_date=payload.start_date,
        end_date=payload.end_date,
        notes=payload.notes,
    )
    db.add(ri)
    await db.commit()
    await db.refresh(ri)
    return ri


@router.put("/recurring/{rid}", response_model=RecurringIncomeOut)
async def update_recurring_income(
    rid: int,
    payload: RecurringIncomeUpdate,
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[AsyncSession, Depends(get_db)],
):
    result = await db.execute(select(RecurringIncome).where(RecurringIncome.id == rid, RecurringIncome.user_id == current_user.id))
    ri = result.scalar_one_or_none()
    if not ri:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=_NOT_FOUND_DETAIL)

    for attr in (
        "title",
        "amount",
        "currency",
        "category",
        "recurrence",
        "months",
        "day_of_month",
        "start_date",
        "end_date",
        "notes",
    ):
        v = getattr(payload, attr)
        if v is not None:
            setattr(ri, attr, v)

    await db.commit()
    await db.refresh(ri)
    return ri


@router.delete("/recurring/{rid}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_recurring_income(
    rid: int,
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[AsyncSession, Depends(get_db)],
):
    result = await db.execute(select(RecurringIncome).where(RecurringIncome.id == rid, RecurringIncome.user_id == current_user.id))
    ri = result.scalar_one_or_none()
    if not ri:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=_NOT_FOUND_DETAIL)
    await db.delete(ri)
    await db.commit()


@router.get("/dashboard", response_model=IncomeDashboard)
async def get_income_dashboard(
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[AsyncSession, Depends(get_db)],
    year: Annotated[int, Query(ge=2020, le=2100)],
    month: Annotated[int, Query(ge=1, le=12)],
):
    """Gelir özet paneli: gerçekleşen + tahmini metrikler.

    - this_month_actual: incomes(bu ay)
    - ytd_actual: incomes(yıl başı..son gün dahil)
    - this_month_recurring: bu ay aktif olan recurring_incomes toplamı
    - ytd_recurring: yıl başı..bu ay (dahil) geçen recurring
    - remaining_year_recurring: bu aydan sonraki ay..yıl sonu recurring
    - year_total_estimate: ytd_actual + remaining_year_recurring
    """
    first_day_year = date_type(year, 1, 1)
    first_day_month = date_type(year, month, 1)
    last_day_month = date_type(year, month, calendar.monthrange(year, month)[1])

    # Gerçekleşen — amount_tl (işlem-anı kuruyla sabit) toplamı.
    actual_month_q = await db.execute(
        select(func.coalesce(func.sum(Income.amount_tl), 0)).where(
            Income.user_id == current_user.id,
            Income.date >= first_day_month,
            Income.date <= last_day_month,
        )
    )
    actual_ytd_q = await db.execute(
        select(func.coalesce(func.sum(Income.amount_tl), 0)).where(
            Income.user_id == current_user.id,
            Income.date >= first_day_year,
            Income.date <= last_day_month,
        )
    )
    this_month_actual = Decimal(actual_month_q.scalar_one())
    ytd_actual = Decimal(actual_ytd_q.scalar_one())

    # Periyodik (yıl içi 12 ay tarama). Tahmin → GÜNCEL kurla TL'ye çevrilir
    # (hibrit kur). Kur haritası tek sefer çekilir (döngüde tekrar TCMB yok).
    rec_q = await db.execute(select(RecurringIncome).where(RecurringIncome.user_id == current_user.id))
    recurring = rec_q.scalars().all()
    rates = await currency_svc.fetch_rates() if recurring else {}

    this_month_recurring = Decimal(0)
    ytd_recurring = Decimal(0)
    remaining_year_recurring = Decimal(0)
    for m in range(1, 13):
        for ri in recurring:
            if not _applies_in_month(ri, year, m):
                continue
            amt = currency_svc.convert_to_tl(Decimal(ri.amount), ri.currency or "TRY", rates)
            if m == month:
                this_month_recurring += amt
            if m <= month:
                ytd_recurring += amt
            else:
                remaining_year_recurring += amt

    year_total_estimate = ytd_actual + remaining_year_recurring

    return IncomeDashboard(
        year=year,
        month=month,
        this_month_actual=this_month_actual,
        ytd_actual=ytd_actual,
        this_month_recurring=this_month_recurring,
        ytd_recurring=ytd_recurring,
        remaining_year_recurring=remaining_year_recurring,
        year_total_estimate=year_total_estimate,
    )


# ---------------------------------------------------------------------------
# Realize endpoint'leri — periyodik kayıtları gerçekleşmiş incomes'a aktar
# ---------------------------------------------------------------------------
# Recurring kategori → income kategorisi (recurring'de "sale" yok, gerisi 1:1)
_RECURRING_TO_INCOME_CAT: dict[str, str] = {
    "salary": "salary",
    "rental": "rental",
    "dividend": "dividend",
    "bonus": "bonus",
    "freelance": "freelance",
    "other": "other",
}


def _date_for_period(ri: RecurringIncome, year: int, month: int) -> date_type:
    """Recurring'in o ay-yıl için 'gerçekleştiği gün' tarihini döner (ortak util)."""
    return recurrence.date_for_period(ri, year, month)


async def _realize_one(
    db: AsyncSession,
    ri: RecurringIncome,
    year: int,
    month: int,
    user_id,
    today: date_type | None = None,
) -> int | None:
    """Tek bir periyodik kayıt için verilen ay-yıl income oluşturur.

    Skip durumları (None döner):
    - Periyot bu ay-yılı kapsamıyor (_applies_in_month False)
    - Hedef tarih (year, month, day_of_month) BUGÜNDEN İLERİDE — ödeme günü
      gelmediği için future-dated kayıt yaratılmaz
    - Aynı (recurring, date) kombinasyonu için kayıt zaten var
    """
    if not _applies_in_month(ri, year, month):
        return None
    target_date = _date_for_period(ri, year, month)
    if today is None:
        today = datetime.now(_ISTANBUL).date()
    # Ödeme günü henüz gelmediyse atla — future-dated income yaratma
    if target_date > today:
        return None
    # Mevcut realize var mı? (unique constraint zaten engelliyor; integrity hatasını
    # önceden yakalamak için kontrol)
    existing_q = await db.execute(
        select(Income.id).where(
            Income.user_id == user_id,
            Income.recurring_income_id == ri.id,
            Income.date == target_date,
        )
    )
    if existing_q.scalar_one_or_none() is not None:
        return None
    # Para birimini recurring'den taşı; realize anı kuruyla amount_tl SABİTLE.
    ri_currency = ri.currency or "TRY"
    amount_tl, exchange_rate = await _resolve_amount_tl(ri.amount, ri_currency)
    inc = Income(
        user_id=user_id,
        amount=ri.amount,
        currency=ri_currency,
        amount_tl=amount_tl,
        exchange_rate=exchange_rate,
        category=_RECURRING_TO_INCOME_CAT.get(ri.category, "other"),
        date=target_date,
        description=ri.title,
        recurring_income_id=ri.id,
    )
    db.add(inc)
    await db.flush()  # ID üret
    return inc.id


@router.post("/recurring/{rid}/realize", response_model=RealizeResult)
async def realize_recurring_period(
    rid: int,
    payload: RealizeMonthRequest,
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[AsyncSession, Depends(get_db)],
):
    """Periyodik kaydın belirli bir ay-yılı için income oluştur.
    Idempotent: aynı dönem ikinci kez çağrılırsa skip."""
    result = await db.execute(
        select(RecurringIncome).where(
            RecurringIncome.id == rid,
            RecurringIncome.user_id == current_user.id,
        )
    )
    ri = result.scalar_one_or_none()
    if not ri:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=_RECURRING_NOT_FOUND)
    if not _applies_in_month(ri, payload.year, payload.month):
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="Bu kayıt belirtilen ay-yılında geçerli değil (periyot dışı)",
        )
    today = datetime.now(_ISTANBUL).date()
    target_date = _date_for_period(ri, payload.year, payload.month)
    if target_date > today:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=f"Ödeme günü ({target_date.isoformat()}) henüz gelmedi — gerçekleşti olarak işaretlenemez",
        )
    new_id = await _realize_one(db, ri, payload.year, payload.month, current_user.id, today=today)
    await db.commit()
    if new_id is None:
        return RealizeResult(realized=0, skipped=1, income_ids=[])
    return RealizeResult(realized=1, skipped=0, income_ids=[new_id])


@router.post("/recurring/{rid}/realize-past", response_model=RealizeResult)
async def realize_recurring_past(
    rid: int,
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[AsyncSession, Depends(get_db)],
):
    """Periyodik kaydın start_date'ten bugüne kadar olan tüm geçmiş dönemleri
    income'a aktar. Mevcut realize'ler skip."""
    result = await db.execute(
        select(RecurringIncome).where(
            RecurringIncome.id == rid,
            RecurringIncome.user_id == current_user.id,
        )
    )
    ri = result.scalar_one_or_none()
    if not ri:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=_RECURRING_NOT_FOUND)

    today = datetime.now(_ISTANBUL).date()
    realized_ids: list[int] = []
    skipped = 0
    # Start'tan bugüne kadar her ayı tara (ödeme günü henüz gelmemiş aylar atlanır)
    y, m = ri.start_date.year, ri.start_date.month
    while date_type(y, m, 1) <= today:
        new_id = await _realize_one(db, ri, y, m, current_user.id, today=today)
        if new_id is None:
            # Sadece "applies + ödeme günü geçmiş + zaten var" durumu skipped
            if _applies_in_month(ri, y, m) and _date_for_period(ri, y, m) <= today:
                skipped += 1
        else:
            realized_ids.append(new_id)
        # Sonraki ay
        m += 1
        if m > 12:
            m = 1
            y += 1
    await db.commit()
    return RealizeResult(realized=len(realized_ids), skipped=skipped, income_ids=realized_ids)


async def _get_owned_recurring(rid: int, user: User, db: AsyncSession) -> RecurringIncome:
    result = await db.execute(select(RecurringIncome).where(RecurringIncome.id == rid, RecurringIncome.user_id == user.id))
    ri = result.scalar_one_or_none()
    if not ri:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=_RECURRING_NOT_FOUND)
    return ri


@router.get("/recurring/{rid}/periods", response_model=RecurringPeriodsResult)
async def get_recurring_periods(
    rid: int,
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[AsyncSession, Depends(get_db)],
):
    """Periyodik gelirin geçmiş+güncel dönemlerini gerçekleşme durumuyla döndürür.

    Yanlış işaretlenen realize/skip'i görüp geri almak için (planlı gider ile
    simetrik): her dönem 'pending' | 'realized' (income_id) | 'skipped' (skip_id)."""
    ri = await _get_owned_recurring(rid, current_user, db)
    today = datetime.now(_ISTANBUL).date()

    inc_rows = (
        await db.execute(
            select(Income.date, Income.id).where(
                Income.user_id == current_user.id,
                Income.recurring_income_id == ri.id,
            )
        )
    ).all()
    realized_by_date = dict(inc_rows)

    skip_rows = (
        await db.execute(
            select(RecurringSkip.period_year, RecurringSkip.period_month, RecurringSkip.id).where(
                RecurringSkip.user_id == current_user.id,
                RecurringSkip.kind == "income",
                RecurringSkip.ref_id == ri.id,
            )
        )
    ).all()
    skip_by_period = {(y, m): sid for y, m, sid in skip_rows}

    periods: list[RecurringPeriodStatus] = []
    for y, m, target in recurrence.iter_due_periods(ri, today):
        income_id = realized_by_date.get(target)
        skip_id = skip_by_period.get((y, m))
        if income_id is not None:
            stat = "realized"
        elif skip_id is not None:
            stat = "skipped"
        else:
            stat = "pending"
        periods.append(
            RecurringPeriodStatus(
                year=y,
                month=m,
                target_date=target,
                status=stat,
                income_id=income_id,
                skip_id=skip_id,
            )
        )
    periods.reverse()
    return RecurringPeriodsResult(periods=periods)


@router.post("/recurring/{rid}/unrealize", response_model=RecurringUnrealizeResult)
async def unrealize_recurring_period(
    rid: int,
    payload: RealizeMonthRequest,
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[AsyncSession, Depends(get_db)],
):
    """Bir dönemin gelir realize'ını geri al: o döneme ait gerçek income kaydını sil.

    Idempotent: o dönem zaten realize değilse removed=0."""
    ri = await _get_owned_recurring(rid, current_user, db)
    target_date = _date_for_period(ri, payload.year, payload.month)
    inc = (
        await db.execute(
            select(Income).where(
                Income.user_id == current_user.id,
                Income.recurring_income_id == ri.id,
                Income.date == target_date,
            )
        )
    ).scalar_one_or_none()
    if inc is None:
        return RecurringUnrealizeResult(removed=0)
    await db.delete(inc)
    await db.commit()
    return RecurringUnrealizeResult(removed=1)


@router.post("/recurring/realize-all-past", response_model=RealizeResult)
async def realize_all_recurring_past(
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[AsyncSession, Depends(get_db)],
):
    """Kullanıcının TÜM periyodik kayıtları için bugüne kadar olan tüm
    geçmiş dönemleri income'a aktar."""
    result = await db.execute(select(RecurringIncome).where(RecurringIncome.user_id == current_user.id))
    all_ri = result.scalars().all()

    today = datetime.now(_ISTANBUL).date()
    realized_ids: list[int] = []
    skipped = 0
    for ri in all_ri:
        y, m = ri.start_date.year, ri.start_date.month
        while date_type(y, m, 1) <= today:
            new_id = await _realize_one(db, ri, y, m, current_user.id, today=today)
            if new_id is None:
                if _applies_in_month(ri, y, m) and _date_for_period(ri, y, m) <= today:
                    skipped += 1
            else:
                realized_ids.append(new_id)
            m += 1
            if m > 12:
                m = 1
                y += 1
    await db.commit()
    return RealizeResult(realized=len(realized_ids), skipped=skipped, income_ids=realized_ids)
