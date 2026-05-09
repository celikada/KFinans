import asyncio
import io
import logging
from decimal import Decimal
from typing import Annotated
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, status
from fastapi.responses import StreamingResponse
from sqlalchemy import select, delete
from sqlalchemy.ext.asyncio import AsyncSession
from app.core.deps import get_db, get_current_user
from app.models.stock import StockHolding as StockHoldingModel
from app.models.user import User
from app.schemas.stocks import StockHolding, StockPositionOut
from app.services.aggregator import fetch_gbp_to_usd, fetch_usd_to_tl
from app.services.stocks import fetch_stock_quotes


def convert_to_tl(price: Decimal, currency: str, usd_tl: Decimal, gbp_usd: Decimal) -> Decimal:
    """Hisse senedi fiyatini TL'ye dönüştürür.

    GBp (pence): pence -> GBP -> USD -> TL zinciri uygulanir.
    USD: USD -> TL.
    TRY: direkt.
    Diger: USD varsayimiyla USD -> TL.
    """
    if currency == "TRY":
        return price
    if currency == "GBp":
        gbp = price / Decimal("100")
        usd = gbp * gbp_usd
        return (usd * usd_tl).quantize(Decimal("0.0001"))
    return (price * usd_tl).quantize(Decimal("0.0001"))

logger = logging.getLogger(__name__)
router = APIRouter(prefix="/portfolio/stocks", tags=["stocks"])


@router.get("/holdings", response_model=list[StockHolding])
async def get_stock_holdings(
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    result = await db.execute(
        select(StockHoldingModel).where(StockHoldingModel.user_id == current_user.id)
    )
    rows = result.scalars().all()
    return [
        StockHolding(
            ticker=r.ticker,
            quantity=float(r.quantity),
            name=r.name,
            avg_cost_tl=float(r.avg_cost_tl) if r.avg_cost_tl is not None else None,
            distributor=r.distributor,
        )
        for r in rows
    ]


@router.put("/holdings", response_model=list[StockHolding])
async def save_stock_holdings(
    holdings: list[StockHolding],
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    await db.execute(delete(StockHoldingModel).where(StockHoldingModel.user_id == current_user.id))
    for h in holdings:
        db.add(StockHoldingModel(
            user_id=current_user.id,
            ticker=h.ticker.upper(),
            quantity=h.quantity,
            name=h.name,
            avg_cost_tl=h.avg_cost_tl,
            distributor=h.distributor,
        ))
    await db.commit()
    return holdings


@router.post("/preview", response_model=list[StockPositionOut])
async def stock_preview(
    holdings: list[StockHolding],
    _: Annotated[User, Depends(get_current_user)],
):
    tickers = [h.ticker.upper() for h in holdings]
    quotes, usd_tl, gbp_usd = await asyncio.gather(
        fetch_stock_quotes(tickers),
        fetch_usd_to_tl(),
        fetch_gbp_to_usd(),
    )

    out = []
    for h in holdings:
        ticker = h.ticker.upper()
        q = quotes.get(ticker)
        if not q:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail=f"{ticker} için fiyat alınamadı. Yahoo Finance ticker'ını kontrol edin (ör. THYAO.IS, AAPL).",
            )
        price_tl = convert_to_tl(q.price, q.currency, usd_tl, gbp_usd)

        qty = Decimal(str(h.quantity))
        total_value_tl = (qty * price_tl).quantize(Decimal("0.01"))

        if h.avg_cost_tl is not None:
            avg_cost = Decimal(str(h.avg_cost_tl))
            cost_basis = (qty * avg_cost).quantize(Decimal("0.01"))
            gain_loss = (total_value_tl - cost_basis).quantize(Decimal("0.01"))
            gain_loss_pct = float(gain_loss / cost_basis * 100) if cost_basis > 0 else None
        else:
            avg_cost = cost_basis = gain_loss = None
            gain_loss_pct = None

        out.append(StockPositionOut(
            ticker=ticker,
            name=h.name or q.name,
            quantity=qty,
            currency=q.currency,
            unit_price_original=q.price,
            unit_price_tl=price_tl,
            total_value_tl=total_value_tl,
            avg_cost_tl=avg_cost,
            cost_basis_tl=cost_basis,
            gain_loss_tl=gain_loss,
            gain_loss_pct=gain_loss_pct,
            distributor=h.distributor,
            # FIN-004 (FAZ H): stale flag UI rozet icin
            is_stale=q.is_stale,
            market_state=q.market_state,
        ))
    return out


@router.get("/export")
async def export_stock_holdings(
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    result = await db.execute(
        select(StockHoldingModel).where(StockHoldingModel.user_id == current_user.id)
    )
    rows = result.scalars().all()

    quotes: dict = {}
    if rows:
        try:
            tickers = [r.ticker for r in rows]
            usd_tl, gbp_usd, raw_quotes = await asyncio.gather(
                fetch_usd_to_tl(),
                fetch_gbp_to_usd(),
                fetch_stock_quotes(tickers),
            )
            for ticker, q in raw_quotes.items():
                if q:
                    price_tl = convert_to_tl(q.price, q.currency, usd_tl, gbp_usd).quantize(Decimal("0.01"))
                    quotes[ticker] = (price_tl, q.currency)
        except Exception:
            logger.warning("Hisse export: canlı fiyat alınamadı")

    wb = Workbook()
    ws = wb.active
    ws.title = "Hisse Senedi"
    headers = ["Ticker", "Adet", "İsim", "Birim Fiyat (₺)", "Toplam Değer (₺)", "Ort. Maliyet (₺)", "Kâr/Zarar (₺)", "Kurum"]
    header_fill = PatternFill("solid", fgColor="059669")
    header_font = Font(bold=True, color="FFFFFF")
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for row_idx, holding in enumerate(rows, 2):
        price_tl, _ = quotes.get(holding.ticker, (None, None))
        total = float(holding.quantity) * float(price_tl) if price_tl else None
        avg_cost = float(holding.avg_cost_tl) if holding.avg_cost_tl is not None else None
        gain_loss = None
        if total is not None and avg_cost is not None:
            cost_basis = float(holding.quantity) * avg_cost
            gain_loss = round(total - cost_basis, 2)
        ws.cell(row=row_idx, column=1, value=holding.ticker)
        ws.cell(row=row_idx, column=2, value=float(holding.quantity))
        ws.cell(row=row_idx, column=3, value=holding.name)
        ws.cell(row=row_idx, column=4, value=float(price_tl) if price_tl else "")
        ws.cell(row=row_idx, column=5, value=total if total is not None else "")
        ws.cell(row=row_idx, column=6, value=avg_cost if avg_cost is not None else "")
        ws.cell(row=row_idx, column=7, value=gain_loss if gain_loss is not None else "")
        ws.cell(row=row_idx, column=8, value=holding.distributor or "")

    for col, width in zip("ABCDEFGH", [12, 14, 30, 18, 18, 18, 18, 20]):
        ws.column_dimensions[col].width = width

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=hisse-senedi.xlsx"},
    )


@router.post("/import-mkk", response_model=list[StockHolding])
async def import_stocks_mkk(
    file: UploadFile = File(...),
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """MKK e-Yatırımcı 'Tüm Kıymetler' raporundan hisse senetlerini içe aktarır.

    Filtre:
      Kıymet Sınıfı = 'HS' AND Ek Tanım = 'A' (aktif tradeable pozisyonlar)
    Mapping:
      Üye → distributor, Menkul Kıymet Kodu → ticker (.IS suffix ekleniyor),
      Adet → quantity, Fiyat (TL) → avg_cost_tl, Kıymet Adı → name
    """
    import xlrd

    if not file.filename or not file.filename.lower().endswith((".xls", ".xlsx")):
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="Sadece .xls veya .xlsx dosyası kabul edilir",
        )
    content = await file.read()
    try:
        wb = xlrd.open_workbook(file_contents=content)
        sh = wb.sheet_by_index(0)
    except Exception:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="MKK Excel dosyası okunamadı (.xls binary formatında olmalı)",
        )

    header_row = None
    for r in range(min(sh.nrows, 20)):
        if str(sh.cell_value(r, 0)).strip().lower() == "üye":
            header_row = r
            break
    if header_row is None:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="MKK formatı tanınmadı: 'Üye' başlık satırı bulunamadı",
        )

    parsed: list[StockHolding] = []
    for r in range(header_row + 1, sh.nrows):
        kind = str(sh.cell_value(r, 2)).strip()
        ek_tanim = str(sh.cell_value(r, 5)).strip().upper()
        if kind != "HS" or ek_tanim != "A":
            continue
        code_raw = sh.cell_value(r, 3)
        bist_code = str(code_raw).strip().upper() if code_raw else ""
        if not bist_code:
            continue
        # MKK BIST kodlarını Yahoo Finance ticker formatına çevir
        ticker = bist_code if "." in bist_code else f"{bist_code}.IS"
        name_raw = sh.cell_value(r, 4)
        name = str(name_raw).strip() if name_raw else ""
        try:
            qty = float(sh.cell_value(r, 7))
            price_raw = sh.cell_value(r, 8)
            price = float(price_raw) if price_raw not in ("", "-") else 0
        except (TypeError, ValueError):
            continue
        if qty <= 0:
            continue
        member = str(sh.cell_value(r, 0)).strip() or None
        if member:
            member = member[:50]
        parsed.append(StockHolding(
            ticker=ticker,
            quantity=qty,
            name=name,
            avg_cost_tl=price if price > 0 else None,
            distributor=member,
        ))

    if not parsed:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="Dosyada 'HS' kıymet sınıfında ve 'A' ek tanımında geçerli kayıt bulunamadı",
        )

    await db.execute(delete(StockHoldingModel).where(StockHoldingModel.user_id == current_user.id))
    for h in parsed:
        db.add(StockHoldingModel(
            user_id=current_user.id,
            ticker=h.ticker,
            quantity=h.quantity,
            name=h.name,
            avg_cost_tl=h.avg_cost_tl,
            distributor=h.distributor,
        ))
    await db.commit()

    # Snapshot tetikle ki Finansal Hedef + History güncel kalsın
    try:
        from app.services.snapshot import compute_and_save_snapshot
        await compute_and_save_snapshot(current_user.id, db)
    except Exception as exc:
        logger.warning("MKK Stocks import sonrası snapshot alınamadı: %s", exc)

    return parsed


@router.post("/import", response_model=list[StockHolding])
async def import_stock_holdings(
    file: UploadFile = File(...),
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    from openpyxl import load_workbook

    if not file.filename or not file.filename.endswith((".xlsx", ".xls")):
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="Sadece .xlsx dosyası kabul edilir")

    content = await file.read()
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
    except Exception:
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="Dosya okunamadı")

    parsed: list[StockHolding] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        ticker = str(row[0]).strip().upper() if row[0] else ""
        qty_raw = row[1]
        name = str(row[2]).strip() if len(row) > 2 and row[2] else ""
        avg_cost_raw = row[3] if len(row) > 3 else None
        distributor_raw = row[7] if len(row) > 7 else None
        if not ticker or ticker == "NONE":
            continue
        try:
            qty = float(qty_raw)
        except (TypeError, ValueError):
            continue
        if qty <= 0:
            continue
        avg_cost_tl: float | None = None
        if avg_cost_raw is not None:
            try:
                avg_cost_tl = float(avg_cost_raw)
                if avg_cost_tl <= 0:
                    avg_cost_tl = None
            except (TypeError, ValueError):
                avg_cost_tl = None
        distributor: str | None = None
        if distributor_raw:
            distributor = str(distributor_raw).strip()[:50] or None
        parsed.append(StockHolding(
            ticker=ticker, quantity=qty, name=name,
            avg_cost_tl=avg_cost_tl, distributor=distributor,
        ))

    if not parsed:
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="Geçerli holding bulunamadı")

    await db.execute(delete(StockHoldingModel).where(StockHoldingModel.user_id == current_user.id))
    for h in parsed:
        db.add(StockHoldingModel(
            user_id=current_user.id,
            ticker=h.ticker,
            quantity=h.quantity,
            name=h.name,
            avg_cost_tl=h.avg_cost_tl,
            distributor=h.distributor,
        ))
    await db.commit()
    return parsed
