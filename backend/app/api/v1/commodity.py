"""Kıymetli maden (altın/gümüş) CRUD endpoint'leri."""

import io
import logging
from decimal import Decimal, InvalidOperation
from typing import Annotated

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile, status
from fastapi.responses import StreamingResponse
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.deps import get_current_user, get_db
from app.core.upload_validation import validate_excel_upload
from app.models.commodity import CommodityHolding
from app.models.user import User
from app.schemas.commodity import (
    CommodityCreate,
    CommodityOut,
    CommodityPositionOut,
    CommoditySummaryOut,
    CommodityUpdate,
)
from app.services.commodity import (
    BIGA_GRAM_WEIGHTS,
    BIGA_METAL,
    COIN_GRAM_WEIGHTS,
    calculate_holding_value,
    fetch_metal_prices,
)

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/portfolio/commodities", tags=["commodities"])

_VALID_UNIT_TYPES = {"gram", "biga", "coin"}
_VALID_METALS = {"gold", "silver"}
_VALID_BIGA = frozenset(BIGA_GRAM_WEIGHTS.keys())
_VALID_COIN = frozenset(COIN_GRAM_WEIGHTS.keys())


@router.get("")
async def list_commodities(
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[AsyncSession, Depends(get_db)],
) -> CommoditySummaryOut:
    """Kullanıcının tüm kıymetli maden varlıklarını anlık fiyatlarla döndürür."""
    result = await db.execute(select(CommodityHolding).where(CommodityHolding.user_id == current_user.id).order_by(CommodityHolding.created_at))
    holdings = result.scalars().all()

    prices = await fetch_metal_prices()
    gold_price = prices["gold"]
    silver_price = prices["silver"]
    gold_price_available = gold_price > 0
    silver_price_available = silver_price > 0

    positions: list[CommodityPositionOut] = []
    total_gold_gram = Decimal("0")
    total_silver_gram = Decimal("0")
    total_value_tl = Decimal("0")

    for h in holdings:
        try:
            val = calculate_holding_value(h, gold_price, silver_price)
        except ValueError as exc:
            logger.error("Holding %d değer hesaplanamadı: %s", h.id, exc)
            continue

        gram_eq = val["gram_equivalent"]
        value_tl = val["total_value_tl"]

        if h.metal == "gold":
            total_gold_gram += gram_eq
        else:
            total_silver_gram += gram_eq
        total_value_tl += value_tl

        positions.append(
            CommodityPositionOut(
                id=h.id,
                unit_type=h.unit_type,
                metal=h.metal,
                biga_code=h.biga_code,
                coin_type=h.coin_type,
                quantity=h.quantity,
                notes=h.notes,
                created_at=h.created_at,
                gram_equivalent=gram_eq,
                total_value_tl=value_tl,
                gold_price_tl=gold_price,
                silver_price_tl=silver_price,
            )
        )

    return CommoditySummaryOut(
        positions=positions,
        total_gold_gram=total_gold_gram.quantize(Decimal("0.0001")),
        total_silver_gram=total_silver_gram.quantize(Decimal("0.0001")),
        total_value_tl=total_value_tl.quantize(Decimal("0.01")),
        gold_price_tl=gold_price,
        silver_price_tl=silver_price,
        gold_price_available=gold_price_available,
        silver_price_available=silver_price_available,
    )


@router.post("", response_model=CommodityOut, status_code=status.HTTP_201_CREATED)
async def create_commodity(
    payload: CommodityCreate,
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[AsyncSession, Depends(get_db)],
) -> CommodityHolding:
    """Yeni kıymetli maden varlığı ekler."""
    holding = CommodityHolding(
        user_id=current_user.id,
        unit_type=payload.unit_type,
        metal=payload.metal,
        biga_code=payload.biga_code,
        coin_type=payload.coin_type,
        quantity=payload.quantity,
        notes=payload.notes.strip() if payload.notes else None,
    )
    db.add(holding)
    await db.commit()
    await db.refresh(holding)
    return holding


@router.put("/{holding_id}", response_model=CommodityOut)
async def update_commodity(
    holding_id: int,
    payload: CommodityUpdate,
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[AsyncSession, Depends(get_db)],
) -> CommodityHolding:
    """Kıymetli maden varlığının miktar ve notunu günceller."""
    result = await db.execute(
        select(CommodityHolding).where(
            CommodityHolding.id == holding_id,
            CommodityHolding.user_id == current_user.id,
        )
    )
    holding = result.scalar_one_or_none()
    if not holding:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Kayıt bulunamadı")

    if payload.quantity is not None:
        holding.quantity = payload.quantity
    if payload.notes is not None:
        holding.notes = payload.notes.strip() or None

    await db.commit()
    await db.refresh(holding)
    return holding


@router.delete("/{holding_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_commodity(
    holding_id: int,
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[AsyncSession, Depends(get_db)],
) -> None:
    """Kıymetli maden varlığını siler."""
    result = await db.execute(
        select(CommodityHolding).where(
            CommodityHolding.id == holding_id,
            CommodityHolding.user_id == current_user.id,
        )
    )
    holding = result.scalar_one_or_none()
    if not holding:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Kayıt bulunamadı")
    await db.delete(holding)
    await db.commit()


@router.get(
    "/export",
    responses={500: {"description": "openpyxl kütüphanesi bulunamadı"}},
)
async def export_commodities(
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[AsyncSession, Depends(get_db)],
):
    """Kıymetli maden varlıklarını Excel dosyası olarak indir."""
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError as exc:
        raise HTTPException(status_code=500, detail="openpyxl kütüphanesi bulunamadı") from exc

    result = await db.execute(select(CommodityHolding).where(CommodityHolding.user_id == current_user.id).order_by(CommodityHolding.created_at))
    holdings = result.scalars().all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Altın-Gümüş"

    # Başlık satırı: altın sarısı arka plan, beyaz bold
    header_fill = PatternFill(start_color="D97706", end_color="D97706", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    headers = ["Tür", "Metal", "BiGA Kodu", "Sikke Türü", "Miktar", "Not", "Oluşturma Tarihi"]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for row_idx, h in enumerate(holdings, start=2):
        ws.cell(row=row_idx, column=1, value=h.unit_type)
        ws.cell(row=row_idx, column=2, value=h.metal)
        ws.cell(row=row_idx, column=3, value=h.biga_code or "")
        ws.cell(row=row_idx, column=4, value=h.coin_type or "")
        ws.cell(row=row_idx, column=5, value=float(h.quantity))
        ws.cell(row=row_idx, column=6, value=h.notes or "")
        ws.cell(row=row_idx, column=7, value=str(h.created_at.date()))

    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 16
    ws.column_dimensions["E"].width = 12
    ws.column_dimensions["F"].width = 36
    ws.column_dimensions["G"].width = 18

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=altin-gumus.xlsx"},
    )


def _cell(row: tuple, idx: int):
    """Tuple'dan güvenli hücre erişimi; index yoksa veya değer falsy ise None."""
    return row[idx] if len(row) > idx and row[idx] else None


def _parse_commodity_row(row: tuple, user_id) -> CommodityHolding | None:
    """Excel satırını doğrular ve CommodityHolding üretir; geçersizse None döner."""
    if not row or all(v is None for v in row):
        return None

    unit_type_val = str(row[0]).strip().lower() if row[0] is not None else ""
    if unit_type_val not in _VALID_UNIT_TYPES:
        return None

    metal_cell = _cell(row, 1)
    metal_val = str(metal_cell).strip().lower() if metal_cell is not None else "gold"
    if metal_val not in _VALID_METALS:
        metal_val = "gold"

    quantity = _parse_quantity(row[4] if len(row) > 4 else None)
    if quantity is None:
        return None

    biga_cell = _cell(row, 2)
    coin_cell = _cell(row, 3)
    notes_cell = _cell(row, 5)
    biga_code_val = str(biga_cell).strip().upper() if biga_cell else None
    coin_type_val = str(coin_cell).strip().lower() if coin_cell else None
    notes_val = str(notes_cell).strip() if notes_cell else None

    resolved = _resolve_metal_type(unit_type_val, metal_val, biga_code_val, coin_type_val)
    if resolved is None:
        return None
    resolved_metal, resolved_biga, resolved_coin = resolved

    return CommodityHolding(
        user_id=user_id,
        unit_type=unit_type_val,
        metal=resolved_metal,
        biga_code=resolved_biga,
        coin_type=resolved_coin,
        quantity=quantity,
        notes=notes_val or None,
    )


def _parse_quantity(quantity_val) -> Decimal | None:
    """Miktarı Decimal'e çevirir; geçersiz veya <=0 ise None döner."""
    try:
        quantity = Decimal(str(quantity_val)).quantize(Decimal("0.0001"))
    except (InvalidOperation, TypeError):
        return None
    if quantity <= 0:
        return None
    return quantity


def _resolve_metal_type(
    unit_type_val: str,
    metal_val: str,
    biga_code_val: str | None,
    coin_type_val: str | None,
) -> tuple[str, str | None, str | None] | None:
    """Tür bazlı doğrulama; (metal, biga_code, coin_type) veya geçersizse None."""
    if unit_type_val == "biga":
        if not biga_code_val or biga_code_val not in _VALID_BIGA:
            return None
        return BIGA_METAL[biga_code_val], biga_code_val, None
    if unit_type_val == "coin":
        if not coin_type_val or coin_type_val not in _VALID_COIN:
            return None
        return "gold", None, coin_type_val
    return metal_val, None, None


@router.post(
    "/import",
    response_model=list[CommodityOut],
    status_code=status.HTTP_201_CREATED,
    responses={
        400: {"description": "Geçersiz Excel dosyası"},
        500: {"description": "openpyxl kütüphanesi bulunamadı"},
    },
)
async def import_commodities(
    file: Annotated[UploadFile, File(...)],
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[AsyncSession, Depends(get_db)],
):
    """Excel dosyasından kıymetli maden varlığı içe aktar (append — mevcut kayıtlar silinmez)."""
    try:
        import openpyxl
    except ImportError as exc:
        raise HTTPException(status_code=500, detail="openpyxl kütüphanesi bulunamadı") from exc

    # SEC-009 (FAZ H): magic-byte + boyut + extension dogrulamasi
    content = await validate_excel_upload(file)
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=400, detail="Geçersiz Excel dosyası") from exc

    ws = wb.active
    added: list[CommodityHolding] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        holding = _parse_commodity_row(row, current_user.id)
        if holding is not None:
            db.add(holding)
            added.append(holding)

    await db.commit()
    for h in added:
        await db.refresh(h)

    return added
