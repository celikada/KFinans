import io
import logging
from decimal import Decimal, InvalidOperation

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile, status
from fastapi.responses import StreamingResponse
from sqlalchemy import delete, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.deps import get_current_user, get_db
from app.core.upload_validation import validate_excel_upload
from app.models.bes import BesHolding as BesHoldingModel
from app.models.user import User
from app.schemas.bes import BesHolding

logger = logging.getLogger(__name__)
router = APIRouter(prefix="/portfolio/bes", tags=["bes"])


@router.get("/holdings", response_model=list[BesHolding])
async def get_bes_holdings(
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    result = await db.execute(
        select(BesHoldingModel).where(BesHoldingModel.user_id == current_user.id)
    )
    rows = result.scalars().all()
    return [
        BesHolding(
            plan_name=r.plan_name,
            contract_number=r.contract_number,
            paid_principal=r.paid_principal,
            paid_returns=r.paid_returns,
            govt_contribution=r.govt_contribution,
            govt_returns=r.govt_returns,
        )
        for r in rows
    ]


@router.put("/holdings", response_model=list[BesHolding])
async def save_bes_holdings(
    holdings: list[BesHolding],
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    await db.execute(delete(BesHoldingModel).where(BesHoldingModel.user_id == current_user.id))
    for h in holdings:
        db.add(BesHoldingModel(
            user_id=current_user.id,
            plan_name=h.plan_name.strip(),
            contract_number=h.contract_number.strip() if h.contract_number else None,
            paid_principal=h.paid_principal,
            paid_returns=h.paid_returns,
            govt_contribution=h.govt_contribution,
            govt_returns=h.govt_returns,
        ))
    await db.commit()
    return holdings


@router.get("/export")
async def export_bes_holdings(
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    result = await db.execute(
        select(BesHoldingModel).where(BesHoldingModel.user_id == current_user.id)
    )
    rows = result.scalars().all()

    wb = Workbook()
    ws = wb.active
    ws.title = "BES Holdingleri"
    headers = [
        "Plan Adı",
        "Sözleşme No",
        "Yatırılan (₺)",
        "Yatırım Getirisi (₺)",
        "Devlet Katkısı (₺)",
        "Devlet Katkı Getirisi (₺)",
        "Toplam (₺)",
    ]
    header_fill = PatternFill("solid", fgColor="047857")
    header_font = Font(bold=True, color="FFFFFF")
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for row_idx, holding in enumerate(rows, 2):
        ws.cell(row=row_idx, column=1, value=holding.plan_name)
        ws.cell(row=row_idx, column=2, value=holding.contract_number or "")
        ws.cell(row=row_idx, column=3, value=float(holding.paid_principal))
        ws.cell(row=row_idx, column=4, value=float(holding.paid_returns))
        ws.cell(row=row_idx, column=5, value=float(holding.govt_contribution))
        ws.cell(row=row_idx, column=6, value=float(holding.govt_returns))
        ws.cell(row=row_idx, column=7, value=float(holding.total_value_tl))

    for col, width in zip("ABCDEFG", [30, 18, 16, 18, 18, 22, 16]):
        ws.column_dimensions[col].width = width

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=bes-holdingleri.xlsx"},
    )


def _parse_decimal(value) -> Decimal:
    """Excel cell'i Decimal'e cevir; None/bos/hatali -> 0."""
    if value is None or value == "":
        return Decimal("0")
    try:
        return Decimal(str(value))
    except (InvalidOperation, ValueError, TypeError):
        return Decimal("0")


@router.post("/import", response_model=list[BesHolding])
async def import_bes_holdings(
    file: UploadFile = File(...),
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    from openpyxl import load_workbook

    # SEC-009 (FAZ H): magic-byte + boyut + extension dogrulamasi
    content = await validate_excel_upload(file)
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
    except Exception:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="Dosya okunamadı",
        )

    parsed: list[BesHolding] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        plan_name = str(row[0]).strip() if row[0] else ""
        if not plan_name or plan_name.lower() == "none":
            continue
        contract = str(row[1]).strip() if len(row) > 1 and row[1] else None
        if contract and contract.lower() == "none":
            contract = None
        paid_principal = _parse_decimal(row[2] if len(row) > 2 else None)
        paid_returns = _parse_decimal(row[3] if len(row) > 3 else None)
        govt_contribution = _parse_decimal(row[4] if len(row) > 4 else None)
        govt_returns = _parse_decimal(row[5] if len(row) > 5 else None)

        # En az bir sayisal alan > 0 olmali (tum sifirsa atla — bos satir)
        if paid_principal + paid_returns + govt_contribution + govt_returns <= 0:
            continue

        parsed.append(BesHolding(
            plan_name=plan_name,
            contract_number=contract,
            paid_principal=paid_principal,
            paid_returns=paid_returns,
            govt_contribution=govt_contribution,
            govt_returns=govt_returns,
        ))

    if not parsed:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="Geçerli BES kaydı bulunamadı",
        )

    await db.execute(delete(BesHoldingModel).where(BesHoldingModel.user_id == current_user.id))
    for h in parsed:
        db.add(BesHoldingModel(
            user_id=current_user.id,
            plan_name=h.plan_name,
            contract_number=h.contract_number,
            paid_principal=h.paid_principal,
            paid_returns=h.paid_returns,
            govt_contribution=h.govt_contribution,
            govt_returns=h.govt_returns,
        ))
    await db.commit()
    return parsed
