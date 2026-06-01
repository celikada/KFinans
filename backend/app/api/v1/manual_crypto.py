"""API'siz borsa hesapları için manuel kripto holding CRUD + Excel + preview."""

import logging
from decimal import Decimal
from io import BytesIO
from typing import Annotated

import openpyxl
from fastapi import APIRouter, Depends, File, HTTPException, UploadFile, status
from fastapi.responses import StreamingResponse
from sqlalchemy import delete, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.deps import get_current_user, get_db
from app.core.upload_validation import validate_excel_upload
from app.models.manual_crypto import ManualCryptoHolding
from app.models.user import User
from app.schemas.manual_crypto import (
    ManualCryptoCreate,
    ManualCryptoOut,
    ManualCryptoPositionOut,
    ManualCryptoSummaryOut,
    ManualCryptoUpdate,
)
from app.services.aggregator import (
    fetch_coingecko_prices_by_ids,
    fetch_combined_prices,
    fetch_usd_to_tl,
    lookup_usd_price,
)
from app.services.tefas import fetch_tefas_prices_by_codes

logger = logging.getLogger(__name__)
router = APIRouter(prefix="/manual-crypto", tags=["manual-crypto"])

# linked commodity ID (XAU/XAG) → fetch_metal_prices() anahtari (gold/silver)
_COMMODITY_METAL_KEYS = {"XAU": "gold", "XAG": "silver"}

_HOLDING_NOT_FOUND = "Manuel kripto kaydı bulunamadı"


def _calc_gain_loss(quantity: Decimal, total_tl: Decimal, avg_cost_tl: Decimal | None) -> tuple[Decimal | None, Decimal | None, float | None]:
    """Avg cost varsa: cost_basis, gain_loss, gain_loss_pct döner; yoksa None'lar."""
    if avg_cost_tl is None or avg_cost_tl <= 0:
        return None, None, None
    cost_basis = (quantity * avg_cost_tl).quantize(Decimal("0.01"))
    gain_loss = (total_tl - cost_basis).quantize(Decimal("0.01"))
    pct = float(gain_loss / cost_basis * 100) if cost_basis > 0 else None
    return cost_basis, gain_loss, pct


class _PriceContext:
    """_enrich_positions icinde toplu cekilen fiyat kaynaklarini tasiyan kap."""

    def __init__(
        self,
        usd_tl: Decimal,
        prices: dict,
        metal_prices: dict[str, Decimal],
        cg_prices_by_id: dict[str, Decimal],
        tefas_prices: dict[str, Decimal],
    ) -> None:
        self.usd_tl = usd_tl
        self.prices = prices
        self.metal_prices = metal_prices
        self.cg_prices_by_id = cg_prices_by_id
        self.tefas_prices = tefas_prices


def _usd_from_tl(unit_tl: Decimal, usd_tl: Decimal) -> Decimal:
    """TL fiyattan USD karsiligi (kur > 0 ve fiyat > 0 ise)."""
    if usd_tl > 0 and unit_tl > 0:
        return (unit_tl / usd_tl).quantize(Decimal("0.000001"))
    return Decimal(0)


def _tl_from_usd(usd: Decimal, usd_tl: Decimal) -> Decimal:
    """USD fiyattan TL karsiligi."""
    return (usd * usd_tl).quantize(Decimal("0.0001")) if usd > 0 else Decimal(0)


def _price_linked(h: ManualCryptoHolding, ctx: _PriceContext) -> tuple[Decimal, Decimal]:
    """linked_source bazinda (usd, unit_tl) cozumle."""
    ls = h.linked_source
    lid = h.linked_id or ""
    if ls == "commodity":
        key = _COMMODITY_METAL_KEYS.get(lid.upper())
        if not key:
            return Decimal(0), Decimal(0)
        unit_tl = ctx.metal_prices.get(key, Decimal(0))
        return _usd_from_tl(unit_tl, ctx.usd_tl), unit_tl
    if ls == "binance":
        usd = lookup_usd_price(lid, ctx.prices)
        return usd, _tl_from_usd(usd, ctx.usd_tl)
    if ls == "coingecko":
        usd = ctx.cg_prices_by_id.get(lid, Decimal(0))
        return usd, _tl_from_usd(usd, ctx.usd_tl)
    if ls == "tefas":
        unit_tl = ctx.tefas_prices.get(lid, Decimal(0))
        return _usd_from_tl(unit_tl, ctx.usd_tl), unit_tl
    return Decimal(0), Decimal(0)


def _resolve_unit_price(h: ManualCryptoHolding, ctx: _PriceContext) -> tuple[Decimal, Decimal]:
    """price_source bazinda (usd, unit_tl) doner."""
    if h.price_source == "manual":
        unit_tl = h.manual_unit_price_tl or Decimal(0)
        return _usd_from_tl(unit_tl, ctx.usd_tl), unit_tl
    if h.price_source == "linked":
        return _price_linked(h, ctx)
    # auto
    usd = lookup_usd_price(h.symbol, ctx.prices)
    return usd, _tl_from_usd(usd, ctx.usd_tl)


def _build_position(h: ManualCryptoHolding, usd: Decimal, unit_tl: Decimal) -> tuple[ManualCryptoPositionOut, Decimal]:
    """Bir holding icin pozisyon cikti'si + TL deger doner."""
    value_tl = (h.quantity * unit_tl).quantize(Decimal("0.01"))
    cost_basis, gain_loss, gain_loss_pct = _calc_gain_loss(h.quantity, value_tl, h.avg_cost_tl)
    position = ManualCryptoPositionOut(
        id=h.id,
        exchange=h.exchange,
        label=h.label,
        symbol=h.symbol,
        quantity=h.quantity,
        avg_cost_tl=h.avg_cost_tl,
        price_source=h.price_source,
        manual_unit_price_tl=h.manual_unit_price_tl,
        linked_source=h.linked_source,
        linked_id=h.linked_id,
        unit_price_usd=usd,
        unit_price_tl=unit_tl,
        total_value_tl=value_tl,
        cost_basis_tl=cost_basis,
        gain_loss_tl=gain_loss,
        gain_loss_pct=gain_loss_pct,
        notes=h.notes,
    )
    return position, value_tl


async def _enrich_positions(
    holdings: list[ManualCryptoHolding],
) -> ManualCryptoSummaryOut:
    """price_source bazlı 3 yol:
    - 'auto'   → Binance USDT + CoinGecko fallback (sembol bazlı)
    - 'manual' → kullanıcının manual_unit_price_tl
    - 'linked' → linked_source bazında dispatch:
                 commodity (XAU/XAG), binance:SYMBOL, coingecko:ID, tefas:CODE
    """
    if not holdings:
        return ManualCryptoSummaryOut(positions=[], total_value_tl=Decimal(0), unknown_symbols=[])

    ctx = await _build_price_context(holdings)

    positions: list[ManualCryptoPositionOut] = []
    total_tl = Decimal(0)
    unknown: list[str] = []

    for h in holdings:
        usd, unit_tl = _resolve_unit_price(h, ctx)
        position, value_tl = _build_position(h, usd, unit_tl)
        total_tl += value_tl
        if unit_tl <= 0:
            unknown.append(h.symbol)
        positions.append(position)

    return ManualCryptoSummaryOut(
        positions=positions,
        total_value_tl=total_tl.quantize(Decimal("0.01")),
        unknown_symbols=sorted(set(unknown)),
    )


async def _build_price_context(holdings: list[ManualCryptoHolding]) -> _PriceContext:
    """Tum holding'ler icin gereken fiyat kaynaklarini bir kerede toplar."""
    auto_holdings = [h for h in holdings if h.price_source == "auto"]
    linked_holdings = [h for h in holdings if h.price_source == "linked"]

    # Hangi kaynaklara hangi ID'ler için sorgu lazım — bir kerede topla
    auto_symbols = list({h.symbol for h in auto_holdings})
    binance_linked_symbols = [h.linked_id for h in linked_holdings if h.linked_source == "binance" and h.linked_id]
    cg_linked_ids = list({h.linked_id for h in linked_holdings if h.linked_source == "coingecko" and h.linked_id})
    tefas_linked_codes = list({h.linked_id for h in linked_holdings if h.linked_source == "tefas" and h.linked_id})
    needs_commodity = any(h.linked_source == "commodity" for h in linked_holdings)

    # USD/TL her durumda lazım
    binance_symbols_to_fetch = list(set(auto_symbols + binance_linked_symbols))
    if binance_symbols_to_fetch:
        prices, usd_tl = await _fetch_prices_safe(binance_symbols_to_fetch)
    else:
        prices = {}
        usd_tl = await fetch_usd_to_tl()

    metal_prices: dict[str, Decimal] = {}
    if needs_commodity:
        try:
            from app.services.commodity import fetch_metal_prices

            metal_prices = await fetch_metal_prices()  # {gold, silver} TRY/g
        except Exception:
            metal_prices = {}

    cg_prices_by_id: dict[str, Decimal] = await fetch_coingecko_prices_by_ids(cg_linked_ids) if cg_linked_ids else {}
    tefas_prices: dict[str, Decimal] = await fetch_tefas_prices_by_codes(tefas_linked_codes) if tefas_linked_codes else {}

    return _PriceContext(
        usd_tl=usd_tl,
        prices=prices,
        metal_prices=metal_prices,
        cg_prices_by_id=cg_prices_by_id,
        tefas_prices=tefas_prices,
    )


async def _fetch_prices_safe(symbols: list[str]) -> tuple[dict[str, Decimal], Decimal]:
    """Binance + CoinGecko fallback + USD/TL — hata durumunda 503.

    SEC-007 (FAZ H): Exception detail'i client'a sizdirilmaz; full trace ops log'a.
    """
    try:
        prices = await fetch_combined_prices(symbols)
    except Exception:
        logger.exception("fetch_combined_prices failed (manual_crypto)")
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="Fiyat bilgisi şu an alınamıyor. Lütfen daha sonra tekrar deneyin.",
        )
    try:
        usd_tl = await fetch_usd_to_tl()
    except Exception:
        logger.exception("fetch_usd_to_tl failed (manual_crypto)")
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="Döviz kuru şu an alınamıyor. Lütfen daha sonra tekrar deneyin.",
        )
    return prices, usd_tl


@router.get("", response_model=ManualCryptoSummaryOut)
async def list_manual_crypto(
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[AsyncSession, Depends(get_db)],
):
    """Tüm manuel kripto pozisyonları + anlık fiyatla TL değer + kâr/zarar."""
    rows = (
        (
            await db.execute(
                select(ManualCryptoHolding)
                .where(ManualCryptoHolding.user_id == current_user.id)
                .order_by(ManualCryptoHolding.exchange, ManualCryptoHolding.symbol)
            )
        )
        .scalars()
        .all()
    )
    return await _enrich_positions(list(rows))


@router.post("", response_model=ManualCryptoOut, status_code=status.HTTP_201_CREATED)
async def create_manual_crypto(
    payload: ManualCryptoCreate,
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[AsyncSession, Depends(get_db)],
):
    holding = ManualCryptoHolding(
        user_id=current_user.id,
        exchange=payload.exchange,
        label=payload.label,
        symbol=payload.symbol,
        quantity=payload.quantity,
        avg_cost_tl=payload.avg_cost_tl,
        price_source=payload.price_source,
        manual_unit_price_tl=payload.manual_unit_price_tl if payload.price_source == "manual" else None,
        linked_source=payload.linked_source if payload.price_source == "linked" else None,
        linked_id=payload.linked_id if payload.price_source == "linked" else None,
        notes=payload.notes,
    )
    db.add(holding)
    await db.commit()
    await db.refresh(holding)
    return holding


@router.put(
    "/{holding_id}",
    response_model=ManualCryptoOut,
    responses={404: {"description": "Manuel kripto kaydı bulunamadı"}},
)
async def update_manual_crypto(
    holding_id: int,
    payload: ManualCryptoUpdate,
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[AsyncSession, Depends(get_db)],
):
    holding = (
        await db.execute(
            select(ManualCryptoHolding).where(
                ManualCryptoHolding.id == holding_id,
                ManualCryptoHolding.user_id == current_user.id,
            )
        )
    ).scalar_one_or_none()
    if not holding:
        raise HTTPException(status_code=404, detail=_HOLDING_NOT_FOUND)

    # Duz alanlar — None olmayan degerler dogrudan yazilir
    for attr in ("exchange", "label", "symbol", "quantity", "avg_cost_tl", "notes"):
        value = getattr(payload, attr)
        if value is not None:
            setattr(holding, attr, value)

    if payload.price_source is not None:
        holding.price_source = payload.price_source
        # mod değişince ilgisiz alanları temizle
        if payload.price_source != "manual":
            holding.manual_unit_price_tl = None
        if payload.price_source != "linked":
            holding.linked_source = None
            holding.linked_id = None

    if payload.manual_unit_price_tl is not None and holding.price_source == "manual":
        holding.manual_unit_price_tl = payload.manual_unit_price_tl
    if holding.price_source == "linked":
        if payload.linked_source is not None:
            holding.linked_source = payload.linked_source
        if payload.linked_id is not None:
            holding.linked_id = payload.linked_id

    await db.commit()
    await db.refresh(holding)
    return holding


@router.delete(
    "/{holding_id}",
    status_code=status.HTTP_204_NO_CONTENT,
    responses={404: {"description": "Manuel kripto kaydı bulunamadı"}},
)
async def delete_manual_crypto(
    holding_id: int,
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[AsyncSession, Depends(get_db)],
):
    result = await db.execute(
        delete(ManualCryptoHolding).where(
            ManualCryptoHolding.id == holding_id,
            ManualCryptoHolding.user_id == current_user.id,
        )
    )
    if result.rowcount == 0:
        raise HTTPException(status_code=404, detail=_HOLDING_NOT_FOUND)
    await db.commit()


@router.get("/export")
async def export_manual_crypto(
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[AsyncSession, Depends(get_db)],
):
    """Excel export — kullanıcı dosyayı düzenleyip import edebilir."""
    rows = (
        (
            await db.execute(
                select(ManualCryptoHolding)
                .where(ManualCryptoHolding.user_id == current_user.id)
                .order_by(ManualCryptoHolding.exchange, ManualCryptoHolding.symbol)
            )
        )
        .scalars()
        .all()
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Manuel Kripto"
    headers = [
        "Borsa",
        "Etiket",
        "Sembol",
        "Miktar",
        "Ort. Maliyet (TL)",
        "Fiyat Kaynagi",
        "Manuel Fiyat (TL)",
        "Linked Source",
        "Linked ID",
        "Notlar",
    ]
    ws.append(headers)
    # Header stilini biraz belirginleştir
    for cell in ws[1]:
        cell.font = openpyxl.styles.Font(bold=True, color="FFFFFF")
        cell.fill = openpyxl.styles.PatternFill(start_color="F59E0B", end_color="F59E0B", fill_type="solid")

    for r in rows:
        ws.append(
            [
                r.exchange,
                r.label or "",
                r.symbol,
                float(r.quantity),
                float(r.avg_cost_tl) if r.avg_cost_tl else "",
                r.price_source,
                float(r.manual_unit_price_tl) if r.manual_unit_price_tl else "",
                r.linked_source or "",
                r.linked_id or "",
                r.notes or "",
            ]
        )

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=manuel_kripto.xlsx"},
    )


def _avg_cost_or_none(raw) -> Decimal | None:
    """Excel hucresinden pozitif Decimal; bos/0/negatif -> None.

    Hatali sayi (InvalidOperation vb.) caller'in outer try/except'ine
    propagate eder — orijinal davranis korunur (satir error listesine duser).
    """
    if raw in (None, ""):
        return None
    value = Decimal(str(raw))
    return value if value > 0 else None


def _manual_price_or_none(raw) -> Decimal | None:
    """Manuel fiyat hucresi — bos/0/negatif/hatali tum durumlarda None (sessiz)."""
    if raw in (None, ""):
        return None
    try:
        value = Decimal(str(raw))
    except Exception:
        return None
    return value if value > 0 else None


def _cell(row, idx: int) -> str | None:
    """row[idx] -> strip edilmis string; index yoksa veya bos ise None."""
    if idx >= len(row) or not row[idx]:
        return None
    return str(row[idx]).strip()


def _resolve_linked_import(price_source: str, source_raw: str | None, id_raw: str | None) -> tuple[str | None, str | None]:
    """Import satirinda linked_source/linked_id cozumle (gecersizse None/None)."""
    if price_source == "linked" and source_raw in ("binance", "coingecko", "tefas", "commodity") and id_raw:
        return source_raw, id_raw[:100]
    return None, None


def _parse_import_row(row, idx: int, user_id, errors: list[str]) -> ManualCryptoHolding | None:
    """Bir Excel satirini ManualCryptoHolding'e cevir.

    Zorunlu alan eksik/gecersizse errors'a mesaj ekler ve None doner.
    """
    exchange = _cell(row, 0) or ""
    label = _cell(row, 1)
    symbol_raw = _cell(row, 2)
    symbol = symbol_raw.upper() if symbol_raw else ""
    qty_raw = row[3]
    avg_cost_raw = row[4] if len(row) > 4 else None
    price_src_cell = _cell(row, 5)
    price_src_raw = price_src_cell.lower() if price_src_cell else "auto"
    manual_price_raw = row[6] if len(row) > 6 else None
    linked_source_cell = _cell(row, 7)
    linked_source_raw = linked_source_cell.lower() if linked_source_cell else None
    linked_id_raw = _cell(row, 8)
    notes = _cell(row, 9)

    if not exchange or not symbol or qty_raw in (None, ""):
        errors.append(f"Satır {idx}: Borsa, Sembol ve Miktar zorunlu")
        return None
    quantity = Decimal(str(qty_raw))
    if quantity <= 0:
        errors.append(f"Satır {idx}: Miktar pozitif olmalı")
        return None

    avg_cost = _avg_cost_or_none(avg_cost_raw)
    price_source = price_src_raw if price_src_raw in ("auto", "manual", "linked") else "auto"
    manual_price = _manual_price_or_none(manual_price_raw) if price_source == "manual" else None
    linked_source, linked_id = _resolve_linked_import(price_source, linked_source_raw, linked_id_raw)

    return ManualCryptoHolding(
        user_id=user_id,
        exchange=exchange[:40],
        label=label[:100] if label else None,
        symbol=symbol[:20],
        quantity=quantity,
        avg_cost_tl=avg_cost,
        price_source=price_source,
        manual_unit_price_tl=manual_price,
        linked_source=linked_source,
        linked_id=linked_id,
        notes=notes[:500] if notes else None,
    )


@router.post("/import")
async def import_manual_crypto(
    file: Annotated[UploadFile, File()],
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[AsyncSession, Depends(get_db)],
):
    """Excel import. Mevcut tüm manuel kayıtlar SİLİNİP yenisi yüklenir
    (replace-all semantik — expenses/income import ile aynı pattern)."""
    # SEC-009 (FAZ H): magic-byte + boyut + extension dogrulamasi
    content = await validate_excel_upload(file)
    try:
        wb = openpyxl.load_workbook(BytesIO(content), data_only=True)
    except Exception:
        # SEC-007 (FAZ H): openpyxl InvalidFileException / BadZipFile mesaji
        # dosya yapisi/path hakkinda ipucu verebilir. Sadece generic mesaj.
        logger.exception("Excel parse failed (manual_crypto import)")
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="Excel dosyası okunamadı. Lütfen .xlsx formatında geçerli bir dosya yükleyin.",
        )
    ws = wb.active
    if ws.max_row < 2:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="Boş dosya — en az 1 satır veri gerekli.",
        )

    new_rows: list[ManualCryptoHolding] = []
    errors: list[str] = []
    for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row or all(v in (None, "") for v in row):
            continue
        try:
            holding = _parse_import_row(row, idx, current_user.id, errors)
            if holding is not None:
                new_rows.append(holding)
        except Exception as e:
            errors.append(f"Satır {idx}: {e}")

    if errors and not new_rows:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="Hiçbir satır işlenemedi: " + "; ".join(errors[:5]),
        )

    # Replace-all — mevcut kayıtları sil, yenilerini ekle
    await db.execute(delete(ManualCryptoHolding).where(ManualCryptoHolding.user_id == current_user.id))
    db.add_all(new_rows)
    await db.commit()
    return {"imported": len(new_rows), "errors": errors}
