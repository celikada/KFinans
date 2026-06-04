import calendar
import io
import logging
from collections.abc import Callable
from datetime import date as date_type
from decimal import Decimal, InvalidOperation
from typing import Annotated

from fastapi import APIRouter, Depends, File, HTTPException, Query, Request, UploadFile, status
from fastapi.responses import StreamingResponse
from sqlalchemy import desc, func, or_, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.config import settings
from app.core.credits import deduct_credits
from app.core.deps import get_current_user, get_db
from app.core.limiter import limiter
from app.core.upload_validation import validate_excel_upload
from app.models.budget import Budget
from app.models.expense import Expense
from app.models.income import Income
from app.models.user import User
from app.schemas.expense import (
    EXPENSE_CATEGORIES,
    CategoryBreakdown,
    ExpenseCreate,
    ExpenseOut,
    ExpenseSummary,
    ExpenseUpdate,
)
from app.schemas.expense_analysis import (
    AnalysisPeriod,
    ExpenseAnalysisOut,
    ExpenseAnalysisRequest,
)
from app.services.audit import AuditAction, log_audit

logger = logging.getLogger(__name__)
router = APIRouter(prefix="/expenses", tags=["expenses"])

# Faz 3 (docs/06 §3): Gelismis harcama analizi 3 kredi tuketir. advice.py'deki
# ADVICE_COST=1 deseninin muadili. Kredi kontrolu LLM cagrisindan ONCE, dusum
# AI basariyla yanit verdikten SONRA (deduct_credits commit ETMEZ, caller commit'ler).
EXPENSE_ANALYSIS_COST = 3

# Kategori key -> insana okunur Turkce etiket (AI prompt veri ozetinde kullanilir).
_CATEGORY_TR_LABEL: dict[str, str] = {
    "food": "Yiyecek",
    "groceries": "Market",
    "transport": "Ulaşım",
    "bills": "Faturalar",
    "health": "Sağlık",
    "entertainment": "Eğlence",
    "clothing": "Giyim",
    "home": "Ev",
    "tax": "Vergi",
    "other": "Diğer",
}

# Türkçe label -> İngilizce key haritası (import için)
LABEL_TO_KEY: dict[str, str] = {
    "yiyecek": "food",
    "market": "groceries",
    "ulaşım": "transport",
    "faturalar": "bills",
    "sağlık": "health",
    "eğlence": "entertainment",
    "giyim": "clothing",
    "ev": "home",
    "vergi": "tax",
    "diğer": "other",
    "food": "food",
    "groceries": "groceries",
    "transport": "transport",
    "bills": "bills",
    "health": "health",
    "entertainment": "entertainment",
    "clothing": "clothing",
    "home": "home",
    "tax": "tax",
    "other": "other",
}


@router.get("", response_model=list[ExpenseOut])
async def list_expenses(
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[AsyncSession, Depends(get_db)],
    year: Annotated[int | None, Query(ge=2020, le=2100)] = None,
    month: Annotated[int | None, Query(ge=1, le=12)] = None,
    category: Annotated[str | None, Query()] = None,
):
    """Harcama listesi — opsiyonel year/month/category filtresi.

    En son tarihliden eskiye sirayla doner; ayni gun icindekiler en son
    eklenen ust sirada.
    """
    stmt = select(Expense).where(Expense.user_id == current_user.id)
    if year is not None and month is not None:
        first_day = date_type(year, month, 1)
        last_day = date_type(year, month, calendar.monthrange(year, month)[1])
        stmt = stmt.where(Expense.date >= first_day, Expense.date <= last_day)
    if category:
        if category not in EXPENSE_CATEGORIES:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail=f"Geçersiz kategori: {category}",
            )
        stmt = stmt.where(Expense.category == category)
    stmt = stmt.order_by(desc(Expense.date), desc(Expense.created_at))

    result = await db.execute(stmt)
    return result.scalars().all()


@router.post("", response_model=ExpenseOut, status_code=status.HTTP_201_CREATED)
async def create_expense(
    payload: ExpenseCreate,
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[AsyncSession, Depends(get_db)],
):
    expense = Expense(
        user_id=current_user.id,
        amount=payload.amount,
        category=payload.category,
        date=payload.date,
        description=payload.description.strip() if payload.description else None,
        credit_card_id=payload.credit_card_id,
        is_paid=payload.is_paid,
    )
    db.add(expense)
    await db.commit()
    await db.refresh(expense)
    return expense


@router.put("/{expense_id}", response_model=ExpenseOut)
async def update_expense(
    expense_id: int,
    payload: ExpenseUpdate,
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[AsyncSession, Depends(get_db)],
):
    result = await db.execute(
        select(Expense).where(
            Expense.id == expense_id,
            Expense.user_id == current_user.id,
        )
    )
    expense = result.scalar_one_or_none()
    if not expense:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Harcama bulunamadı")

    if payload.amount is not None:
        expense.amount = payload.amount
    if payload.category is not None:
        expense.category = payload.category
    if payload.date is not None:
        expense.date = payload.date
    if payload.description is not None:
        expense.description = payload.description.strip() or None
    # credit_card_id: None değeri "kart bağlantısını kaldır" anlamına gelmesi için
    # özel handling — payload'da explicit "credit_card_id" yoksa atla, varsa None bile
    # set edilebilir. Pydantic model_fields_set ile kontrol.
    if "credit_card_id" in payload.model_fields_set:
        expense.credit_card_id = payload.credit_card_id
    if payload.is_paid is not None:
        expense.is_paid = payload.is_paid

    await db.commit()
    await db.refresh(expense)
    return expense


@router.delete("/{expense_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_expense(
    expense_id: int,
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[AsyncSession, Depends(get_db)],
):
    result = await db.execute(
        select(Expense).where(
            Expense.id == expense_id,
            Expense.user_id == current_user.id,
        )
    )
    expense = result.scalar_one_or_none()
    if not expense:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Harcama bulunamadı")

    await db.delete(expense)
    await db.commit()


@router.get("/summary", response_model=ExpenseSummary)
async def get_expense_summary(
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[AsyncSession, Depends(get_db)],
    year: Annotated[int, Query(ge=2020, le=2100)],
    month: Annotated[int, Query(ge=1, le=12)],
):
    """Belirli ay icin toplam + kategori bazinda kirilim.

    Cift sayim kurali: credit_card_id NOT NULL + is_paid=true olan kayitlar
    haric tutulur — bu harcamalar kart borcuyla zaten sayildi (credit_cards
    + statements + installments tarafinda).
    """
    first_day = date_type(year, month, 1)
    last_day = date_type(year, month, calendar.monthrange(year, month)[1])

    # Cift sayim filtresi: kart + odendi olanlari haric tut
    not_double_counted = or_(
        Expense.credit_card_id.is_(None),
        Expense.is_paid.is_(False),
    )

    # Toplam ve adet
    total_q = await db.execute(
        select(func.coalesce(func.sum(Expense.amount), 0), func.count(Expense.id)).where(
            Expense.user_id == current_user.id,
            Expense.date >= first_day,
            Expense.date <= last_day,
            not_double_counted,
        )
    )
    total, count = total_q.one()

    # Kategori kirilimi (ayni filtre)
    cat_q = await db.execute(
        select(
            Expense.category,
            func.sum(Expense.amount),
            func.count(Expense.id),
        )
        .where(
            Expense.user_id == current_user.id,
            Expense.date >= first_day,
            Expense.date <= last_day,
            not_double_counted,
        )
        .group_by(Expense.category)
        .order_by(desc(func.sum(Expense.amount)))
    )
    breakdown = [CategoryBreakdown(category=cat, total=Decimal(amt), count=cnt) for cat, amt, cnt in cat_q.all()]

    return ExpenseSummary(
        year=year,
        month=month,
        total=Decimal(total),
        count=count,
        by_category=breakdown,
    )


@router.get("/export", responses={500: {"description": "openpyxl kütüphanesi bulunamadı"}})
async def export_expenses(
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[AsyncSession, Depends(get_db)],
    year: Annotated[int | None, Query(ge=2020, le=2100)] = None,
    month: Annotated[int | None, Query(ge=1, le=12)] = None,
):
    """Harcamaları Excel dosyası olarak indir."""
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        raise HTTPException(status_code=500, detail="openpyxl kütüphanesi bulunamadı")

    stmt = select(Expense).where(Expense.user_id == current_user.id)
    if year is not None and month is not None:
        first_day = date_type(year, month, 1)
        last_day = date_type(year, month, calendar.monthrange(year, month)[1])
        stmt = stmt.where(Expense.date >= first_day, Expense.date <= last_day)
    stmt = stmt.order_by(Expense.date, Expense.created_at)

    result = await db.execute(stmt)
    expenses = result.scalars().all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Harcamalar"

    # Başlık satırı: kırmızı arka plan, beyaz bold
    header_fill = PatternFill(start_color="DC2626", end_color="DC2626", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    headers = ["Tarih", "Kategori", "Tutar (₺)", "Açıklama"]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    # Veri satırları
    for row_idx, exp in enumerate(expenses, start=2):
        ws.cell(row=row_idx, column=1, value=str(exp.date))
        ws.cell(row=row_idx, column=2, value=exp.category)
        ws.cell(row=row_idx, column=3, value=float(exp.amount))
        ws.cell(row=row_idx, column=4, value=exp.description or "")

    # Kolon genişliği
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 40

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=harcamalar.xlsx"},
    )


def _parse_excel_date(date_val) -> date_type | None:
    """Excel hücresinden tarih parse — desteklenmeyen/boş ise None."""
    if isinstance(date_val, date_type):
        return date_val
    if isinstance(date_val, str):
        try:
            return date_type.fromisoformat(date_val.strip())
        except ValueError:
            return None
    # Excel numeric date
    try:
        from openpyxl.utils.datetime import from_excel

        return from_excel(date_val).date() if date_val is not None else None
    except Exception:
        return None


def _parse_amount(amount_val) -> Decimal | None:
    """Tutar parse — pozitif değilse veya geçersizse None."""
    try:
        amount = Decimal(str(amount_val)).quantize(Decimal("0.01"))
    except (InvalidOperation, TypeError):
        return None
    return amount if amount > 0 else None


def _expense_from_row(row, user_id) -> Expense | None:
    """Bir Excel satırından Expense üretir; geçersiz satırda None döner."""
    n = len(row)
    date_val, cat_val, amount_val, desc_val = (row[i] if i < n else None for i in range(4))

    parsed_date = _parse_excel_date(date_val)
    if parsed_date is None:
        return None

    cat_str = str(cat_val).strip().lower() if cat_val is not None else ""
    category = LABEL_TO_KEY.get(cat_str)
    if not category:
        return None

    amount = _parse_amount(amount_val)
    if amount is None:
        return None

    description = str(desc_val).strip() if desc_val else None
    return Expense(
        user_id=user_id,
        amount=amount,
        category=category,
        date=parsed_date,
        description=description or None,
    )


@router.post(
    "/import",
    response_model=list[ExpenseOut],
    status_code=status.HTTP_201_CREATED,
    responses={
        400: {"description": "Geçersiz Excel dosyası"},
        500: {"description": "openpyxl kütüphanesi bulunamadı"},
    },
)
async def import_expenses(
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[AsyncSession, Depends(get_db)],
    file: Annotated[UploadFile, File()],
):
    """Excel dosyasından harcama içe aktar (append — mevcut kayıtlar silinmez)."""
    try:
        import openpyxl
    except ImportError:
        raise HTTPException(status_code=500, detail="openpyxl kütüphanesi bulunamadı")

    # SEC-009 (FAZ H): magic-byte + boyut + extension dogrulamasi
    content = await validate_excel_upload(file)
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    except Exception:
        raise HTTPException(status_code=400, detail="Geçersiz Excel dosyası")

    ws = wb.active
    added: list[Expense] = []

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or all(v is None for v in row):
            continue
        exp = _expense_from_row(row, current_user.id)
        if exp is None:
            continue
        db.add(exp)
        added.append(exp)

    await db.commit()
    for exp in added:
        await db.refresh(exp)

    return added


def _month_range(months: int) -> tuple[date_type, date_type, list[tuple[int, int]]]:
    """Bu ay dahil son `months` ayin (ilk_gun, son_gun) ve (yil, ay) listesi.

    `months_seq` eskiden yeniye sirali (örn. [(2026,1),(2026,2),...]).
    """
    today = date_type.today()
    months_seq: list[tuple[int, int]] = []
    y, m = today.year, today.month
    for _ in range(months):
        months_seq.append((y, m))
        m -= 1
        if m == 0:
            m = 12
            y -= 1
    months_seq.reverse()  # eskiden yeniye
    first_y, first_m = months_seq[0]
    last_y, last_m = months_seq[-1]
    start = date_type(first_y, first_m, 1)
    end = date_type(last_y, last_m, calendar.monthrange(last_y, last_m)[1])
    return start, end, months_seq


def _monthly_breakdown_lines(
    matrix: dict[tuple[int, int], dict[str, Decimal]],
    months_seq: list[tuple[int, int]],
    fmt: Callable[[Decimal], str],
) -> list[str]:
    """Ay × kategori matrisini '- YYYY-MM (toplam ...): ...' satirlarina cevirir."""
    out: list[str] = []
    for y, m in months_seq:
        cell = matrix.get((y, m), {})
        month_total = sum(cell.values(), Decimal("0"))
        if not cell:
            out.append(f"- {y:04d}-{m:02d}: kayıt yok")
            continue
        parts = ", ".join(f"{_CATEGORY_TR_LABEL.get(c, c)} {fmt(v)} TL" for c, v in sorted(cell.items(), key=lambda kv: kv[1], reverse=True))
        out.append(f"- {y:04d}-{m:02d} (toplam {fmt(month_total)} TL): {parts}")
    return out


async def _build_expense_data_summary(
    db: AsyncSession,
    user_id,
    months: int,
) -> tuple[str, Decimal, int, str, str]:
    """AI'a gonderilecek harcama veri ozetini (duz metin) uretir.

    Doner: (summary_text, total_expense, expense_count, month_from, month_to).
    Cift sayim filtresi (kart + odendi haric) get_expense_summary ile ayni.
    Gelir (incomes) ve butce (budgets) ozetini de — varsa — ekler.
    """
    start, end, months_seq = _month_range(months)
    month_from = f"{months_seq[0][0]:04d}-{months_seq[0][1]:02d}"
    month_to = f"{months_seq[-1][0]:04d}-{months_seq[-1][1]:02d}"

    not_double_counted = or_(
        Expense.credit_card_id.is_(None),
        Expense.is_paid.is_(False),
    )

    # Ay × kategori matrisi (year, month, category -> toplam)
    rows = (
        await db.execute(
            select(
                func.extract("year", Expense.date).label("y"),
                func.extract("month", Expense.date).label("m"),
                Expense.category,
                func.sum(Expense.amount),
            )
            .where(
                Expense.user_id == user_id,
                Expense.date >= start,
                Expense.date <= end,
                not_double_counted,
            )
            .group_by("y", "m", Expense.category)
        )
    ).all()

    matrix: dict[tuple[int, int], dict[str, Decimal]] = {(y, m): {} for (y, m) in months_seq}
    cat_totals: dict[str, Decimal] = {}
    total_expense = Decimal("0")
    for y, m, category, amt in rows:
        key = (int(y), int(m))
        amt_d = Decimal(amt)
        matrix.setdefault(key, {})[category] = amt_d
        cat_totals[category] = cat_totals.get(category, Decimal("0")) + amt_d
        total_expense += amt_d

    # Kayit adedi (filtreli)
    expense_count = (
        await db.execute(
            select(func.count(Expense.id)).where(
                Expense.user_id == user_id,
                Expense.date >= start,
                Expense.date <= end,
                not_double_counted,
            )
        )
    ).scalar_one()

    def _fmt(d: Decimal) -> str:
        # 12345.67 -> "12.345,67" (TR binlik ayraci)
        s = f"{d:,.2f}"
        return s.replace(",", "X").replace(".", ",").replace("X", ".")

    lines: list[str] = []
    lines.append(f"## Dönem\n{months} ay incelendi: {month_from} – {month_to}.")
    lines.append(f"Toplam harcama (çift sayım hariç): {_fmt(total_expense)} TL, {expense_count} kayıt.")

    # Kategori toplamlari (buyukten kucuge) + pay
    lines.append("\n## Kategori toplamları (dönem geneli)")
    if cat_totals:
        for category, ctot in sorted(cat_totals.items(), key=lambda kv: kv[1], reverse=True):
            label = _CATEGORY_TR_LABEL.get(category, category)
            pct = (ctot / total_expense * 100) if total_expense > 0 else Decimal("0")
            lines.append(f"- {label} ({category}): {_fmt(ctot)} TL — %{pct:.1f}")
    else:
        lines.append("- (Bu dönemde harcama kaydı yok.)")

    # Ay × kategori matrisi
    lines.append("\n## Aylık kategori dağılımı")
    lines.extend(_monthly_breakdown_lines(matrix, months_seq, _fmt))

    # Gelir ozeti (opsiyonel — gelir-gider dengesi yorumu icin)
    total_income = (
        await db.execute(
            select(func.coalesce(func.sum(Income.amount), 0)).where(
                Income.user_id == user_id,
                Income.date >= start,
                Income.date <= end,
            )
        )
    ).scalar_one()
    if total_income and total_income > 0:
        lines.append("\n## Gelir (aynı dönem)")
        lines.append(f"- Toplam gelir: {_fmt(Decimal(total_income))} TL")
        net = Decimal(total_income) - total_expense
        lines.append(f"- Net (gelir - gider): {_fmt(net)} TL")

    # Butce hedefleri (opsiyonel — kategori bazli kiyas)
    budgets = (await db.execute(select(Budget.category, Budget.amount).where(Budget.user_id == user_id))).all()
    if budgets:
        lines.append("\n## Aylık bütçe hedefleri (kategori başına)")
        for category, amount in budgets:
            label = _CATEGORY_TR_LABEL.get(category, category)
            lines.append(f"- {label} ({category}): aylık {_fmt(Decimal(amount))} TL hedef")

    return "\n".join(lines), total_expense, int(expense_count), month_from, month_to


@router.post(
    "/analysis/generate",
    response_model=ExpenseAnalysisOut,
    status_code=status.HTTP_201_CREATED,
)
@limiter.limit("5/hour")
async def generate_expense_analysis(
    request: Request,
    payload: ExpenseAnalysisRequest,
    current_user: Annotated[User, Depends(get_current_user)],
    db: Annotated[AsyncSession, Depends(get_db)],
):
    """Faz 3: Harcama AI analizi — son N ayin kategori bazli kirilimini Claude'a
    gonderir, egitsel markdown analiz doner. 3 kredi tuketir.

    advice.py deseniyle birebir: (1) Anthropic acik riza kontrolu (403, KVKK m.9 —
    harcama verisi de yurt disina gidiyor), (2) kredi on-kontrol LLM'den ONCE (402),
    (3) AI cagrisi, (4) basari sonrasi deduct_credits, (5) log_audit, (6) tek commit.
    Sonuc ephemeral — DB'ye yazilmaz (kredi tuketimi credit_transactions'ta izlenir).
    """
    from app.services.expense_analyst import ExpenseAnalystService

    # KVKK m.9: harcama verisi Anthropic'e (yurt disi) aktarildigi icin acik riza sart.
    if current_user.anthropic_consent_at is None:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail=(
                "Anthropic API'ye veri aktarimi icin acik riza gerekli (KVKK m.9). Ayarlar > Gizlilik bolumunden 'Anthropic AI tavsiye' onayini etkinlestirin."
            ),
        )

    # Kredi kontrolu LLM cagrisindan ONCE — Anthropic'i bos cagirmamak icin.
    if (current_user.credit_balance or 0) < EXPENSE_ANALYSIS_COST:
        raise HTTPException(
            status_code=status.HTTP_402_PAYMENT_REQUIRED,
            detail=f"Yetersiz kredi. Harcama analizi basina {EXPENSE_ANALYSIS_COST} kredi gerekir; mevcut: {current_user.credit_balance}.",
        )

    summary_text, total_expense, expense_count, month_from, month_to = await _build_expense_data_summary(
        db,
        current_user.id,
        payload.months,
    )

    analyst = ExpenseAnalystService()
    # AI fail -> exception deduct_credits'ten ONCE firlar -> kredi dusmez.
    analysis = await analyst.analyze_expenses(user=current_user, data_summary=summary_text)

    # Atomik dusum (deduct_credits flush eder, commit caller'da). reference_id yok
    # (ephemeral — kalici advice kaydi olusturulmuyor).
    balance_after = await deduct_credits(
        db,
        user_id=current_user.id,
        amount=EXPENSE_ANALYSIS_COST,
        reason="expense_analysis",
        extra={"months": payload.months, "month_from": month_from, "month_to": month_to},
    )

    await log_audit(
        db,
        request,
        action=AuditAction.EXPENSE_ANALYSIS_GENERATE,
        user_id=current_user.id,
        resource=f"expense_analysis:{month_from}..{month_to}",
        extra={
            "months": payload.months,
            "model": settings.claude_model,
            "credits_used": EXPENSE_ANALYSIS_COST,
            "credit_balance_after": balance_after,
            "expense_count": expense_count,
        },
    )

    await db.commit()

    return ExpenseAnalysisOut(
        analysis=analysis,
        credits_used=EXPENSE_ANALYSIS_COST,
        period=AnalysisPeriod(
            months=payload.months,
            month_from=month_from,
            month_to=month_to,
            total_expense=float(total_expense),
            expense_count=expense_count,
        ),
    )
