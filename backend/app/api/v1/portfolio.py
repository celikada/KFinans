import io
from decimal import Decimal
from typing import Annotated
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, status
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from sqlalchemy import select, desc, delete
from sqlalchemy.orm import selectinload
from sqlalchemy.ext.asyncio import AsyncSession
from app.core.deps import get_db, get_current_user
from app.core.security import decrypt_secret
from app.models.integration import Integration
from app.models.portfolio import PortfolioSnapshot
from app.models.tefas import TefasHolding as TefasHoldingModel
from app.models.user import User
from app.schemas.portfolio import SnapshotOut, PortfolioChanges, PortfolioBreakdown, StakingPosition
from app.services.aggregator import fetch_usd_to_tl
from app.services.exchange.binance import BinanceService
from app.services.exchange.icrypex import ICrypexService

router = APIRouter(prefix="/portfolio", tags=["portfolio"])


@router.get("", response_model=SnapshotOut)
async def get_current_portfolio(
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    result = await db.execute(
        select(PortfolioSnapshot)
        .where(PortfolioSnapshot.user_id == current_user.id)
        .options(selectinload(PortfolioSnapshot.asset_positions))
        .order_by(desc(PortfolioSnapshot.snapshot_date))
        .limit(1)
    )
    snapshot = result.scalar_one_or_none()
    if not snapshot:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Henüz portföy verisi yok")
    return snapshot


@router.get("/history", response_model=list[SnapshotOut])
async def get_portfolio_history(
    limit: int = 12,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    result = await db.execute(
        select(PortfolioSnapshot)
        .where(PortfolioSnapshot.user_id == current_user.id)
        .options(selectinload(PortfolioSnapshot.asset_positions))
        .order_by(desc(PortfolioSnapshot.snapshot_date))
        .limit(limit)
    )
    return result.scalars().all()


@router.get("/changes", response_model=PortfolioChanges)
async def get_portfolio_changes(
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    # Son 5 snapshot yeterli (WoW + MoM için)
    result = await db.execute(
        select(PortfolioSnapshot)
        .where(PortfolioSnapshot.user_id == current_user.id)
        .order_by(desc(PortfolioSnapshot.snapshot_date))
        .limit(5)
    )
    snapshots = result.scalars().all()
    if not snapshots:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Henüz portföy verisi yok")

    from app.services.aggregator import calculate_changes
    return calculate_changes(snapshots)


@router.get("/breakdown", response_model=PortfolioBreakdown)
async def get_portfolio_breakdown(
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    result = await db.execute(
        select(PortfolioSnapshot)
        .where(PortfolioSnapshot.user_id == current_user.id)
        .options(selectinload(PortfolioSnapshot.asset_positions))
        .order_by(desc(PortfolioSnapshot.snapshot_date))
        .limit(1)
    )
    snapshot = result.scalar_one_or_none()
    if not snapshot:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Henüz portföy verisi yok")

    from app.services.aggregator import calculate_breakdown
    return calculate_breakdown(snapshot)


class TefasHolding(BaseModel):
    code: str
    quantity: float
    name: str = ""


class TefasPositionOut(BaseModel):
    code: str
    name: str
    quantity: Decimal
    unit_price_tl: Decimal
    total_value_tl: Decimal


@router.get("/tefas/holdings", response_model=list[TefasHolding])
async def get_tefas_holdings(
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    result = await db.execute(
        select(TefasHoldingModel).where(TefasHoldingModel.user_id == current_user.id)
    )
    rows = result.scalars().all()
    return [TefasHolding(code=r.code, quantity=float(r.quantity), name=r.name) for r in rows]


@router.put("/tefas/holdings", response_model=list[TefasHolding])
async def save_tefas_holdings(
    holdings: list[TefasHolding],
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    await db.execute(delete(TefasHoldingModel).where(TefasHoldingModel.user_id == current_user.id))
    for h in holdings:
        db.add(TefasHoldingModel(user_id=current_user.id, code=h.code.upper(), quantity=h.quantity, name=h.name))
    await db.commit()
    return holdings


@router.get("/tefas/export")
async def export_tefas_holdings(
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from app.services.tefas import TefasService

    result = await db.execute(
        select(TefasHoldingModel).where(TefasHoldingModel.user_id == current_user.id)
    )
    rows = result.scalars().all()

    # Canlı fiyat çek
    prices: dict[str, Decimal] = {}
    if rows:
        try:
            svc = TefasService([{"code": r.code, "quantity": float(r.quantity), "name": r.name} for r in rows])
            assets = await svc.fetch()
            prices = {a.symbol: a.unit_price_tl for a in assets}
        except Exception:
            pass

    wb = Workbook()
    ws = wb.active
    ws.title = "TEFAS Holdingleri"

    # Başlık satırı
    headers = ["Fon Kodu", "Adet", "İsim", "Birim Fiyat (₺)", "Toplam Değer (₺)"]
    header_fill = PatternFill("solid", fgColor="1D4ED8")
    header_font = Font(bold=True, color="FFFFFF")
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    # Veri satırları
    for row_idx, holding in enumerate(rows, 2):
        unit_price = prices.get(holding.code, Decimal("0"))
        total = float(holding.quantity) * float(unit_price)
        ws.cell(row=row_idx, column=1, value=holding.code)
        ws.cell(row=row_idx, column=2, value=float(holding.quantity))
        ws.cell(row=row_idx, column=3, value=holding.name)
        ws.cell(row=row_idx, column=4, value=float(unit_price) if unit_price else "")
        ws.cell(row=row_idx, column=5, value=total if unit_price else "")

    # Sütun genişlikleri
    for col, width in zip("ABCDE", [12, 14, 30, 18, 18]):
        ws.column_dimensions[col].width = width

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=tefas-holdingleri.xlsx"},
    )


@router.post("/tefas/import", response_model=list[TefasHolding])
async def import_tefas_holdings(
    file: UploadFile = File(...),
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    from openpyxl import load_workbook

    if not file.filename or not file.filename.endswith((".xlsx", ".xls")):
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="Sadece .xlsx dosyası kabul edilir")

    content = await file.read()
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
    except Exception:
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="Dosya okunamadı")

    parsed: list[TefasHolding] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        code = str(row[0]).strip().upper() if row[0] else ""
        qty_raw = row[1]
        name = str(row[2]).strip() if len(row) > 2 and row[2] else ""
        if not code or code == "NONE":
            continue
        try:
            qty = float(qty_raw)
        except (TypeError, ValueError):
            continue
        if qty <= 0:
            continue
        parsed.append(TefasHolding(code=code, quantity=qty, name=name))

    if not parsed:
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="Geçerli holding bulunamadı")

    await db.execute(delete(TefasHoldingModel).where(TefasHoldingModel.user_id == current_user.id))
    for h in parsed:
        db.add(TefasHoldingModel(user_id=current_user.id, code=h.code, quantity=h.quantity, name=h.name))
    await db.commit()
    return parsed


@router.post("/tefas/preview", response_model=list[TefasPositionOut])
async def tefas_preview(
    holdings: list[TefasHolding],
    _: Annotated[User, Depends(get_current_user)],
):
    """Girilen fon kodları için TEFAS'tan canlı fiyat çeker, kaydetmez."""
    from app.services.tefas import TefasService
    svc = TefasService([{"code": h.code, "quantity": h.quantity, "name": h.name} for h in holdings])
    try:
        assets = await svc.fetch()
    except ValueError as e:
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail=str(e))
    return [
        TefasPositionOut(
            code=a.symbol,
            name=a.name,
            quantity=a.liquid_quantity,
            unit_price_tl=a.unit_price_tl,
            total_value_tl=a.liquid_quantity * a.unit_price_tl,
        )
        for a in assets
    ]


class CryptoPositionOut(BaseModel):
    provider: str
    symbol: str
    liquid_quantity: Decimal
    staked_quantity: Decimal
    unit_price_usd: Decimal
    unit_price_tl: Decimal
    total_value_tl: Decimal


@router.get("/crypto", response_model=list[CryptoPositionOut])
async def get_crypto_positions(
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    result = await db.execute(
        select(Integration).where(
            Integration.user_id == current_user.id,
            Integration.provider.in_(["binance", "icrypex"]),
            Integration.is_active.is_(True),
        )
    )
    integrations = result.scalars().all()
    if not integrations:
        return []

    usd_tl = await fetch_usd_to_tl()
    all_assets = []

    for intg in integrations:
        try:
            api_key = decrypt_secret(intg.encrypted_key)
            if intg.provider == "binance":
                api_secret = decrypt_secret(intg.encrypted_secret) if intg.encrypted_secret else ""
                svc = BinanceService(api_key, api_secret)
            else:
                svc = ICrypexService(api_key)
            assets = await svc.fetch()
            all_assets.extend(assets)
        except Exception:
            pass

    return [
        CryptoPositionOut(
            provider=a.provider,
            symbol=a.symbol,
            liquid_quantity=a.liquid_quantity,
            staked_quantity=a.staked_quantity,
            unit_price_usd=a.unit_price_usd,
            unit_price_tl=(a.unit_price_usd * usd_tl).quantize(Decimal("0.01")),
            total_value_tl=(
                (a.liquid_quantity + a.staked_quantity) * a.unit_price_usd * usd_tl
            ).quantize(Decimal("0.01")),
        )
        for a in all_assets
    ]


@router.get("/staking", response_model=list[StakingPosition])
async def get_staking_positions(
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    result = await db.execute(
        select(PortfolioSnapshot)
        .where(PortfolioSnapshot.user_id == current_user.id)
        .options(selectinload(PortfolioSnapshot.asset_positions))
        .order_by(desc(PortfolioSnapshot.snapshot_date))
        .limit(1)
    )
    snapshot = result.scalar_one_or_none()
    if not snapshot:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Henüz portföy verisi yok")

    from app.services.aggregator import extract_staking_positions
    return extract_staking_positions(snapshot)
