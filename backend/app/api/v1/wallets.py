import io
import uuid
from typing import Annotated

from fastapi import APIRouter, Depends, File, HTTPException, Request, UploadFile, status
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.deps import get_current_user, get_db
from app.core.limiter import limiter
from app.core.masking import mask_address as _mask_address  # BACK-013: central helper
from app.core.security import verify_password
from app.core.upload_validation import validate_excel_upload
from app.models.integration import WalletAddress
from app.models.user import User
from app.schemas.integration import WalletCreate, WalletOut
from app.services.audit import AuditAction, log_audit

router = APIRouter(prefix="/wallets", tags=["wallets"])

CurrentUser = Annotated[User, Depends(get_current_user)]
DbSession = Annotated[AsyncSession, Depends(get_db)]

VALID_CHAINS = {
    "sonic",
    "avalanche_c",
    "avalanche_p",
    "ethereum",
    "bitcoin",
    "solana",
    "cardano",
    "algorand",
    "polkadot",
    "litecoin",
}


@router.get("", response_model=list[WalletOut])
async def list_wallets(
    current_user: CurrentUser,
    db: DbSession,
):
    result = await db.execute(select(WalletAddress).where(WalletAddress.user_id == current_user.id, WalletAddress.is_active.is_(True)))
    return result.scalars().all()


@router.post("", response_model=WalletOut, status_code=status.HTTP_201_CREATED)
async def add_wallet(
    request: Request,
    payload: WalletCreate,
    current_user: CurrentUser,
    db: DbSession,
):
    if payload.chain not in VALID_CHAINS:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=f"Geçersiz zincir: {payload.chain}",
        )
    wallet = WalletAddress(
        user_id=current_user.id,
        chain=payload.chain,
        address=payload.address,
        label=payload.label,
    )
    db.add(wallet)
    try:
        await db.flush()
    except IntegrityError:
        await db.rollback()
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="Bu cüzdan zaten kayıtlı",
        )
    await log_audit(
        db,
        request,
        action=AuditAction.WALLET_ADD,
        user_id=current_user.id,
        resource=f"wallet:{wallet.id}",
        extra={"chain": payload.chain, "label": payload.label},
    )
    await db.commit()
    await db.refresh(wallet)
    return wallet


@router.delete("/{wallet_id}", status_code=status.HTTP_204_NO_CONTENT)
async def remove_wallet(
    wallet_id: str,
    request: Request,
    current_user: CurrentUser,
    db: DbSession,
):
    result = await db.execute(
        select(WalletAddress).where(
            WalletAddress.id == uuid.UUID(wallet_id),
            WalletAddress.user_id == current_user.id,
        )
    )
    wallet = result.scalar_one_or_none()
    if not wallet:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Cüzdan bulunamadı")
    deleted_chain = wallet.chain
    await db.delete(wallet)
    await log_audit(
        db,
        request,
        action=AuditAction.WALLET_DELETE,
        user_id=current_user.id,
        resource=f"wallet:{wallet_id}",
        extra={"chain": deleted_chain},
    )
    await db.commit()


class WalletExportRequest(BaseModel):
    """Tam-adres export icin sifre dogrulama govdesi."""

    password: str


def _build_wallets_xlsx(rows: list[WalletAddress], include_full_address: bool) -> io.BytesIO:
    """Cüzdan listesinden xlsx üretir. include_full_address=False ise adresler maskelenir."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Blockchain Cüzdanları"

    # Guvenlik uyari satiri (row 1) — Excel'in en ustunde gorunur
    warning_text = (
        "GUVENLIK UYARISI: Bu dosya kripto cuzdan adreslerinizi icerir. "
        "xpub/extended public key sizmasi blockchain bakiye gecmisinizi acik "
        "yapar. Bu dosyayi e-postayla paylasmayin, bulut deposunda sifresiz tutmayin."
    )
    warning_cell = ws.cell(row=1, column=1, value=warning_text)
    warning_cell.font = Font(bold=True, color="C53030")
    warning_cell.fill = PatternFill("solid", fgColor="FED7D7")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=3)

    headers = ["Zincir", "Adres" + ("" if include_full_address else " (maskeli)"), "Etiket"]
    header_fill = PatternFill("solid", fgColor="7C3AED")
    header_font = Font(bold=True, color="FFFFFF")
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=2, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for row_idx, w in enumerate(rows, 3):
        ws.cell(row=row_idx, column=1, value=w.chain)
        addr = w.address if include_full_address else _mask_address(w.address)
        ws.cell(row=row_idx, column=2, value=addr)
        ws.cell(row=row_idx, column=3, value=w.label or "")

    for col, width in zip("ABC", [18, 80 if include_full_address else 30, 20]):
        ws.column_dimensions[col].width = width

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


_XLSX_MEDIA = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_XLSX_HEADERS = {"Content-Disposition": "attachment; filename=blockchain-cuzdanlari.xlsx"}
_EXPORT_RESOURCE = "wallet:export.xlsx"


@router.get("/export")
async def export_wallets(request: Request, current_user: CurrentUser, db: DbSession):
    """COMP-024 (FAZ H): Maskeli adresli Excel export (sifre gerektirmez).

    Adresler maskeli (`xpub6C...4D4D`) — Excel sizmasinda blockchain bakiye
    gecmisi acigi onlenir. **Tam (maskesiz) adres icin POST /wallets/export**
    (kullanici sifresi dogrulanir) kullanin; tam xpub her zaman sifre arkasinda.
    """
    result = await db.execute(select(WalletAddress).where(WalletAddress.user_id == current_user.id, WalletAddress.is_active.is_(True)))
    rows = result.scalars().all()
    buf = _build_wallets_xlsx(rows, include_full_address=False)
    await log_audit(
        db,
        request,
        action=AuditAction.WALLET_EXPORT,
        user_id=current_user.id,
        resource=_EXPORT_RESOURCE,
        extra={"wallet_count": len(rows), "include_full_address": False},
    )
    await db.commit()
    return StreamingResponse(buf, media_type=_XLSX_MEDIA, headers=_XLSX_HEADERS)


@router.post(
    "/export",
    responses={403: {"description": "Şifre hatalı — tam adres verilmez"}},
)
@limiter.limit("5/minute")
async def export_wallets_full(body: WalletExportRequest, request: Request, current_user: CurrentUser, db: DbSession):
    """Tam (maskesiz) adresli Excel export — kullanici sifresi dogrulanir.

    Sifre yanlis/eksikse 403 doner ve tam adres VERILMEZ. Tam xpub indirimi
    audit'lenir (`full=true`) — KVKK m.12 forensic. Rate limit 5/dk (sifre
    brute-force korumasi). xpub sizmasi blockchain bakiye gecmisini acik yapar;
    bu yuzden tam adres her zaman sifre dogrulamasi arkasindadir.
    """
    if not verify_password(body.password, current_user.password_hash):
        await log_audit(
            db,
            request,
            action=AuditAction.WALLET_EXPORT,
            user_id=current_user.id,
            resource=_EXPORT_RESOURCE,
            extra={"full": True, "result": "wrong_password"},
        )
        await db.commit()
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Şifre hatalı — tam adres verilmedi.")

    result = await db.execute(select(WalletAddress).where(WalletAddress.user_id == current_user.id, WalletAddress.is_active.is_(True)))
    rows = result.scalars().all()
    buf = _build_wallets_xlsx(rows, include_full_address=True)
    await log_audit(
        db,
        request,
        action=AuditAction.WALLET_EXPORT,
        user_id=current_user.id,
        resource=_EXPORT_RESOURCE,
        extra={"wallet_count": len(rows), "full": True},
    )
    await db.commit()
    return StreamingResponse(buf, media_type=_XLSX_MEDIA, headers=_XLSX_HEADERS)


def _parse_wallet_row(row: tuple) -> dict | None:
    """Excel satırını doğrular ve cüzdan dict'i üretir; geçersizse None döner."""
    chain = str(row[0]).strip().lower() if row[0] else ""
    if not chain or chain == "none" or chain not in VALID_CHAINS:
        return None

    address = str(row[1]).strip() if len(row) > 1 and row[1] else ""
    if not address or address == "none":
        return None

    label = str(row[2]).strip() if len(row) > 2 and row[2] else None
    if label == "none":
        label = None

    return {"chain": chain, "address": address, "label": label}


@router.post("/import", response_model=list[WalletOut])
async def import_wallets(
    file: Annotated[UploadFile, File()],
    current_user: CurrentUser,
    db: DbSession,
):
    from openpyxl import load_workbook

    # SEC-009 (FAZ H): magic-byte + boyut + extension dogrulamasi
    content = await validate_excel_upload(file)
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
    except Exception as exc:
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="Dosya okunamadı") from exc

    parsed: list[dict] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        wallet_data = _parse_wallet_row(row)
        if wallet_data is not None:
            parsed.append(wallet_data)

    if not parsed:
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="Geçerli cüzdan bulunamadı")

    # Maskeli-adres korumasi: maskeli export (adres "..." icerir) geri import
    # edilirse replace-all gercek tam adresleri ezerdi (veri kaybi). Maskeli satir
    # varsa import'u tamamen reddet — mevcut cuzdanlara DOKUNULMAZ.
    if any("..." in p["address"] for p in parsed):
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=(
                'Bu dosya maskeli adresler ("...") içeriyor; içe aktarılamaz '
                "(mevcut cüzdanlarınız korundu). Tam adresli dosya için “Excel "
                "indir” → şifre doğrulaması ile dışa aktarın."
            ),
        )

    # Mevcut cüzdanları sil, yenilerini ekle
    existing = await db.execute(select(WalletAddress).where(WalletAddress.user_id == current_user.id))
    for w in existing.scalars().all():
        await db.delete(w)
    # DELETE'leri INSERT'lerden önce flush et — aksi halde aynı adres yeniden
    # import edilince (user_id, chain, fingerprint) unique kisiti INSERT'i DELETE'ten
    # once flush ederse 409 verirdi (idempotent re-import).
    await db.flush()

    added = []
    for p in parsed:
        w = WalletAddress(
            user_id=current_user.id,
            chain=p["chain"],
            address=p["address"],
            label=p["label"],
        )
        db.add(w)
        added.append(w)

    await db.commit()
    for w in added:
        await db.refresh(w)
    return added


@router.post("/sync", status_code=status.HTTP_202_ACCEPTED)
async def sync_wallets(current_user: CurrentUser):
    return {"detail": "Blockchain senkronizasyonu başlatıldı"}
