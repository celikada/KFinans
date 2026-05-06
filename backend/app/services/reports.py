"""Rapor üretim servisi: Excel + PDF.

İki tür rapor:
- Cash flow (yıllık nakit akışı): aylar × gerçek/tahmini gelir-gider + net
- Snapshot (portföy özet): varlık listesi + toplamlar

Hem Excel (openpyxl) hem PDF (reportlab) çıktısı üretir.
"""
from decimal import Decimal
from io import BytesIO

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer,
)

# Türkçe karakter desteği için DejaVu Sans (Linux'ta yaygın)
_TR_FONT_REGISTERED = False


def _ensure_tr_font() -> str:
    """Türkçe karakter destekleyen font tescil et. Yoksa Helvetica fallback."""
    global _TR_FONT_REGISTERED
    if _TR_FONT_REGISTERED:
        return "DejaVu"
    try:
        pdfmetrics.registerFont(TTFont("DejaVu", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"))
        pdfmetrics.registerFont(TTFont("DejaVu-Bold", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"))
        _TR_FONT_REGISTERED = True
        return "DejaVu"
    except Exception:
        return "Helvetica"


def _fmt_tl(v) -> str:
    """Decimal/float → '12,345.67 ₺' (Türkçe virgül binlik ayırıcı için TR locale yok,
    o yüzden basit US format)."""
    if v is None:
        return "—"
    n = float(v)
    return f"{n:,.2f} ₺"


# ---------------------------------------------------------------------------
# Cash Flow raporları
# ---------------------------------------------------------------------------
_MONTHS_TR = ["Ocak", "Şubat", "Mart", "Nisan", "Mayıs", "Haziran",
              "Temmuz", "Ağustos", "Eylül", "Ekim", "Kasım", "Aralık"]


def cash_flow_to_xlsx(year: int, months: list, totals: dict) -> bytes:
    """Cash flow yıllık raporunu Excel olarak üret.

    months: list of dicts with keys: month, income_actual, income_forecast,
            expense_actual, expense_forecast, net, is_past
    totals: {total_income, total_expense, total_net}
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Nakit Akışı {year}"

    # Başlık satırı
    headers = [
        "Ay", "Tip", "Gerçek Gelir", "Tahmini Gelir",
        "Gerçek Gider", "Tahmini Gider", "Net",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")

    for m in months:
        ws.append([
            _MONTHS_TR[m["month"] - 1],
            "Geçmiş" if m["is_past"] else "Tahmin",
            float(m["income_actual"]),
            float(m["income_forecast"]),
            float(m["expense_actual"]),
            float(m["expense_forecast"]),
            float(m["net"]),
        ])

    # Toplam satırı
    last_row = ws.max_row + 1
    ws.cell(row=last_row, column=1, value="YIL TOPLAMI").font = Font(bold=True)
    ws.cell(row=last_row, column=2, value="").alignment = Alignment(horizontal="center")
    ws.cell(row=last_row, column=3, value=float(totals["total_income"])).font = Font(bold=True)
    ws.cell(row=last_row, column=5, value=float(totals["total_expense"])).font = Font(bold=True)
    ws.cell(row=last_row, column=7, value=float(totals["total_net"])).font = Font(bold=True)
    for col in range(1, 8):
        ws.cell(row=last_row, column=col).fill = PatternFill(
            start_color="E5E7EB", end_color="E5E7EB", fill_type="solid"
        )

    # Kolon genişlikleri
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 10
    for col in "CDEFG":
        ws.column_dimensions[col].width = 16

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def cash_flow_to_pdf(year: int, months: list, totals: dict) -> bytes:
    """Cash flow yıllık raporunu PDF olarak üret."""
    font = _ensure_tr_font()
    bold_font = font + "-Bold" if font == "DejaVu" else "Helvetica-Bold"

    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4),
                            leftMargin=1.5 * cm, rightMargin=1.5 * cm,
                            topMargin=1.5 * cm, bottomMargin=1.5 * cm)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("Title", parent=styles["Title"], fontName=bold_font, fontSize=18)
    subtitle_style = ParagraphStyle("Sub", parent=styles["Normal"], fontName=font, fontSize=10, textColor=colors.grey)

    elements = []
    elements.append(Paragraph(f"Nakit Akışı Raporu — {year}", title_style))
    elements.append(Paragraph("KFinans · Mayotek", subtitle_style))
    elements.append(Spacer(1, 0.5 * cm))

    # Özet kutusu
    summary_data = [
        ["Toplam Gelir", "Toplam Gider", "Net"],
        [_fmt_tl(totals["total_income"]),
         _fmt_tl(totals["total_expense"]),
         _fmt_tl(totals["total_net"])],
    ]
    summary_table = Table(summary_data, colWidths=[8 * cm, 8 * cm, 8 * cm])
    summary_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F3F4F6")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#6B7280")),
        ("FONTNAME", (0, 0), (-1, 0), bold_font),
        ("FONTNAME", (0, 1), (-1, 1), bold_font),
        ("FONTSIZE", (0, 0), (-1, 0), 10),
        ("FONTSIZE", (0, 1), (-1, 1), 14),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TEXTCOLOR", (0, 1), (0, 1), colors.HexColor("#10B981")),
        ("TEXTCOLOR", (1, 1), (1, 1), colors.HexColor("#EF4444")),
        ("TEXTCOLOR", (2, 1), (2, 1), colors.HexColor("#2563EB")),
        ("TOPPADDING", (0, 0), (-1, -1), 12),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 12),
    ]))
    elements.append(summary_table)
    elements.append(Spacer(1, 0.6 * cm))

    # Aylık tablo
    table_data = [["Ay", "Tip", "Gerçek Gelir", "Tahmini Gelir",
                   "Gerçek Gider", "Tahmini Gider", "Net"]]
    for m in months:
        table_data.append([
            _MONTHS_TR[m["month"] - 1],
            "Geçmiş" if m["is_past"] else "Tahmin",
            _fmt_tl(m["income_actual"]),
            _fmt_tl(m["income_forecast"]),
            _fmt_tl(m["expense_actual"]),
            _fmt_tl(m["expense_forecast"]),
            _fmt_tl(m["net"]),
        ])
    # Toplam satırı
    table_data.append([
        "YIL TOPLAMI", "",
        _fmt_tl(totals["total_income"]), "",
        _fmt_tl(totals["total_expense"]), "",
        _fmt_tl(totals["total_net"]),
    ])

    monthly_table = Table(table_data, repeatRows=1, colWidths=[
        2.5 * cm, 2 * cm, 3.2 * cm, 3.2 * cm, 3.2 * cm, 3.2 * cm, 3.2 * cm,
    ])
    monthly_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2563EB")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), bold_font),
        ("FONTNAME", (0, 1), (-1, -2), font),
        ("FONTNAME", (0, -1), (-1, -1), bold_font),
        ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#E5E7EB")),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("ALIGN", (2, 0), (-1, -1), "RIGHT"),
        ("ALIGN", (0, 0), (1, -1), "LEFT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, colors.HexColor("#F9FAFB")]),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#E5E7EB")),
    ]))
    elements.append(monthly_table)
    elements.append(Spacer(1, 0.4 * cm))

    elements.append(Paragraph(
        "Geçmiş aylar = gerçekleşen değerler (incomes + expenses + ödenmiş ekstreler). "
        "Gelecek aylar = tahmin (recurring incomes + planlı harcamalar + taksitler + henüz ödenmemiş ekstreler). "
        "Çift sayım kuralı uygulandı: kart + ödendi olan harcamalar gider toplamına dahil edilmedi.",
        ParagraphStyle("Footer", parent=styles["Normal"], fontName=font, fontSize=8,
                       textColor=colors.HexColor("#9CA3AF")),
    ))

    doc.build(elements)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Snapshot raporları
# ---------------------------------------------------------------------------
def snapshot_to_xlsx(snapshot, positions: list) -> bytes:
    """Bir snapshot için tüm pozisyon listesini Excel olarak üret."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Snapshot {snapshot.snapshot_date}"

    # Üst panel: tarih + toplam
    ws.cell(row=1, column=1, value="Tarih:").font = Font(bold=True)
    ws.cell(row=1, column=2, value=str(snapshot.snapshot_date))
    ws.cell(row=2, column=1, value="Toplam (TL):").font = Font(bold=True)
    ws.cell(row=2, column=2, value=float(snapshot.total_value_tl))
    if snapshot.usd_try_rate:
        ws.cell(row=3, column=1, value="USD/TRY kuru:").font = Font(bold=True)
        ws.cell(row=3, column=2, value=float(snapshot.usd_try_rate))

    headers = ["Tip", "Provider", "Sembol", "İsim", "Likit", "Stake",
               "Pending", "Birim Fiyat (TL)", "Toplam Değer (TL)", "Ağırlık %"]
    ws.append([])  # boş satır
    ws.append(headers)
    header_row = ws.max_row
    for cell in ws[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="6366F1", end_color="6366F1", fill_type="solid")

    for p in positions:
        ws.append([
            p.asset_type, p.provider, p.symbol, p.name,
            float(p.liquid_quantity),
            float(p.staked_quantity),
            float(p.pending_rewards),
            float(p.unit_price_tl),
            float(p.total_value_tl),
            float(p.weight_pct),
        ])

    # Kolon genişlikleri
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 30
    for col in "EFGHIJ":
        ws.column_dimensions[col].width = 16

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def snapshot_to_pdf(snapshot, positions: list) -> bytes:
    """Bir snapshot için PDF rapor."""
    font = _ensure_tr_font()
    bold_font = font + "-Bold" if font == "DejaVu" else "Helvetica-Bold"

    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4),
                            leftMargin=1.5 * cm, rightMargin=1.5 * cm,
                            topMargin=1.5 * cm, bottomMargin=1.5 * cm)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("Title", parent=styles["Title"], fontName=bold_font, fontSize=18)
    subtitle_style = ParagraphStyle("Sub", parent=styles["Normal"], fontName=font, fontSize=10, textColor=colors.grey)

    elements = []
    elements.append(Paragraph(f"Portföy Snapshot — {snapshot.snapshot_date}", title_style))
    elements.append(Paragraph("KFinans · Mayotek", subtitle_style))
    elements.append(Spacer(1, 0.5 * cm))

    # Özet
    total_tl = float(snapshot.total_value_tl)
    rate = float(snapshot.usd_try_rate) if snapshot.usd_try_rate else None
    total_usd = (total_tl / rate) if rate else None

    summary_data = [
        ["Toplam (TL)", "Toplam (USD)", "USD/TRY", "Pozisyon"],
        [
            _fmt_tl(total_tl),
            f"${total_usd:,.2f}" if total_usd is not None else "—",
            f"{rate:.4f}" if rate else "—",
            str(len(positions)),
        ],
    ]
    summary_table = Table(summary_data, colWidths=[6 * cm, 6 * cm, 6 * cm, 4 * cm])
    summary_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F3F4F6")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#6B7280")),
        ("FONTNAME", (0, 0), (-1, 0), bold_font),
        ("FONTNAME", (0, 1), (-1, 1), bold_font),
        ("FONTSIZE", (0, 0), (-1, 0), 9),
        ("FONTSIZE", (0, 1), (-1, 1), 13),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 10),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 10),
    ]))
    elements.append(summary_table)
    elements.append(Spacer(1, 0.6 * cm))

    # Pozisyon tablosu (en büyük 50; PDF kalabalık olmasın)
    sorted_positions = sorted(positions, key=lambda p: float(p.total_value_tl), reverse=True)[:50]
    table_data = [["Tip", "Provider", "Sembol", "İsim", "Toplam (TL)", "Ağırlık %"]]
    for p in sorted_positions:
        table_data.append([
            p.asset_type, p.provider, p.symbol, p.name[:40],
            _fmt_tl(p.total_value_tl),
            f"%{float(p.weight_pct):.2f}",
        ])

    pos_table = Table(table_data, repeatRows=1, colWidths=[
        2.5 * cm, 3 * cm, 2.5 * cm, 7 * cm, 4 * cm, 3 * cm,
    ])
    pos_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#6366F1")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), bold_font),
        ("FONTNAME", (0, 1), (-1, -1), font),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("ALIGN", (4, 0), (-1, -1), "RIGHT"),
        ("ALIGN", (0, 0), (3, -1), "LEFT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F9FAFB")]),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#E5E7EB")),
    ]))
    elements.append(pos_table)

    if len(positions) > 50:
        elements.append(Spacer(1, 0.3 * cm))
        elements.append(Paragraph(
            f"İlk 50 pozisyon gösteriliyor. Toplam {len(positions)} pozisyon var. Tam liste için Excel raporunu kullanın.",
            ParagraphStyle("Footer", parent=styles["Normal"], fontName=font, fontSize=8,
                           textColor=colors.HexColor("#9CA3AF")),
        ))

    doc.build(elements)
    return buf.getvalue()
