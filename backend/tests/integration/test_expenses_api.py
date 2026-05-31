"""Expenses CRUD + summary endpoint testleri."""

import io

import openpyxl
import pytest
from httpx import AsyncClient

from tests.conftest import make_user


def _exp(amount: float, category: str, date: str, description: str | None = None) -> dict:
    return {"amount": amount, "category": category, "date": date, "description": description}


def _excel_bytes(rows: list[tuple]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(("Tarih", "Kategori", "Tutar", "Açıklama"))
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ─── CRUD ─────────────────────────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_empty_list(client: AsyncClient):
    headers = await make_user(client, "exp_empty@example.com")
    resp = await client.get("/api/v1/expenses", headers=headers)
    assert resp.status_code == 200
    assert resp.json() == []


@pytest.mark.asyncio
async def test_create_expense(client: AsyncClient):
    headers = await make_user(client, "exp_create@example.com")
    resp = await client.post(
        "/api/v1/expenses",
        json=_exp(150.50, "groceries", "2026-05-02", "Migros haftalik"),
        headers=headers,
    )
    assert resp.status_code == 201
    data = resp.json()
    assert float(data["amount"]) == 150.50
    assert data["category"] == "groceries"
    assert data["date"] == "2026-05-02"
    assert data["description"] == "Migros haftalik"
    assert "id" in data


@pytest.mark.asyncio
async def test_create_invalid_category_returns_422(client: AsyncClient):
    headers = await make_user(client, "exp_bad_cat@example.com")
    resp = await client.post(
        "/api/v1/expenses",
        json=_exp(50, "yolo", "2026-05-02"),
        headers=headers,
    )
    assert resp.status_code == 422


@pytest.mark.asyncio
async def test_create_negative_amount_returns_422(client: AsyncClient):
    headers = await make_user(client, "exp_neg@example.com")
    resp = await client.post(
        "/api/v1/expenses",
        json=_exp(-10, "food", "2026-05-02"),
        headers=headers,
    )
    assert resp.status_code == 422


@pytest.mark.asyncio
async def test_create_zero_amount_returns_422(client: AsyncClient):
    """gt=0 — tam sifir kabul edilmez."""
    headers = await make_user(client, "exp_zero@example.com")
    resp = await client.post(
        "/api/v1/expenses",
        json=_exp(0, "food", "2026-05-02"),
        headers=headers,
    )
    assert resp.status_code == 422


@pytest.mark.asyncio
async def test_update_expense(client: AsyncClient):
    headers = await make_user(client, "exp_update@example.com")
    create = await client.post(
        "/api/v1/expenses",
        json=_exp(50, "food", "2026-05-02", "ilk"),
        headers=headers,
    )
    expense_id = create.json()["id"]

    update = await client.put(
        f"/api/v1/expenses/{expense_id}",
        json={"amount": 75.50, "description": "guncellendi"},
        headers=headers,
    )
    assert update.status_code == 200
    assert float(update.json()["amount"]) == 75.50
    assert update.json()["description"] == "guncellendi"
    assert update.json()["category"] == "food"  # degismedi


@pytest.mark.asyncio
async def test_update_nonexistent_returns_404(client: AsyncClient):
    headers = await make_user(client, "exp_404@example.com")
    resp = await client.put(
        "/api/v1/expenses/99999",
        json={"amount": 100},
        headers=headers,
    )
    assert resp.status_code == 404


@pytest.mark.asyncio
async def test_delete_expense(client: AsyncClient):
    headers = await make_user(client, "exp_delete@example.com")
    create = await client.post(
        "/api/v1/expenses",
        json=_exp(50, "food", "2026-05-02"),
        headers=headers,
    )
    expense_id = create.json()["id"]

    delete = await client.delete(f"/api/v1/expenses/{expense_id}", headers=headers)
    assert delete.status_code == 204

    get_after = await client.get("/api/v1/expenses", headers=headers)
    assert get_after.json() == []


@pytest.mark.asyncio
async def test_list_filters_by_year_month(client: AsyncClient):
    headers = await make_user(client, "exp_filter@example.com")
    await client.post("/api/v1/expenses", json=_exp(100, "food", "2026-04-15"), headers=headers)
    await client.post("/api/v1/expenses", json=_exp(200, "food", "2026-05-10"), headers=headers)
    await client.post("/api/v1/expenses", json=_exp(300, "food", "2026-05-20"), headers=headers)

    resp = await client.get("/api/v1/expenses?year=2026&month=5", headers=headers)
    assert resp.status_code == 200
    data = resp.json()
    assert len(data) == 2  # sadece Mayis kayitlari
    amounts = sorted(float(e["amount"]) for e in data)
    assert amounts == [200.0, 300.0]


@pytest.mark.asyncio
async def test_list_filters_by_category(client: AsyncClient):
    headers = await make_user(client, "exp_cat_filter@example.com")
    await client.post("/api/v1/expenses", json=_exp(50, "food", "2026-05-01"), headers=headers)
    await client.post("/api/v1/expenses", json=_exp(100, "transport", "2026-05-02"), headers=headers)
    await client.post("/api/v1/expenses", json=_exp(75, "food", "2026-05-03"), headers=headers)

    resp = await client.get("/api/v1/expenses?category=food", headers=headers)
    data = resp.json()
    assert len(data) == 2
    assert all(e["category"] == "food" for e in data)


@pytest.mark.asyncio
async def test_list_invalid_category_returns_422(client: AsyncClient):
    headers = await make_user(client, "exp_bad_filter@example.com")
    resp = await client.get("/api/v1/expenses?category=invalid", headers=headers)
    assert resp.status_code == 422


@pytest.mark.asyncio
async def test_list_sorted_desc_by_date(client: AsyncClient):
    headers = await make_user(client, "exp_sort@example.com")
    await client.post("/api/v1/expenses", json=_exp(10, "food", "2026-05-01"), headers=headers)
    await client.post("/api/v1/expenses", json=_exp(20, "food", "2026-05-15"), headers=headers)
    await client.post("/api/v1/expenses", json=_exp(30, "food", "2026-05-10"), headers=headers)

    resp = await client.get("/api/v1/expenses", headers=headers)
    dates = [e["date"] for e in resp.json()]
    assert dates == ["2026-05-15", "2026-05-10", "2026-05-01"]


# ─── Summary ──────────────────────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_summary_empty_month(client: AsyncClient):
    headers = await make_user(client, "sum_empty@example.com")
    resp = await client.get("/api/v1/expenses/summary?year=2026&month=5", headers=headers)
    assert resp.status_code == 200
    data = resp.json()
    assert data["year"] == 2026
    assert data["month"] == 5
    assert float(data["total"]) == 0
    assert data["count"] == 0
    assert data["by_category"] == []


@pytest.mark.asyncio
async def test_summary_with_data(client: AsyncClient):
    headers = await make_user(client, "sum_data@example.com")
    await client.post("/api/v1/expenses", json=_exp(100, "food", "2026-05-05"), headers=headers)
    await client.post("/api/v1/expenses", json=_exp(50, "food", "2026-05-10"), headers=headers)
    await client.post("/api/v1/expenses", json=_exp(200, "transport", "2026-05-12"), headers=headers)
    # Farkli ay — summary'ye dahil olmamali
    await client.post("/api/v1/expenses", json=_exp(999, "food", "2026-04-15"), headers=headers)

    resp = await client.get("/api/v1/expenses/summary?year=2026&month=5", headers=headers)
    data = resp.json()
    assert float(data["total"]) == 350.0
    assert data["count"] == 3

    by_cat = {b["category"]: b for b in data["by_category"]}
    assert float(by_cat["food"]["total"]) == 150.0
    assert by_cat["food"]["count"] == 2
    assert float(by_cat["transport"]["total"]) == 200.0
    assert by_cat["transport"]["count"] == 1


# ─── IDOR ─────────────────────────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_user_a_cannot_see_user_b_expenses(client: AsyncClient):
    a = await make_user(client, "exp_idor_a@example.com")
    b = await make_user(client, "exp_idor_b@example.com")
    await client.post("/api/v1/expenses", json=_exp(100, "food", "2026-05-02"), headers=b)

    resp_a = await client.get("/api/v1/expenses", headers=a)
    assert resp_a.json() == []


@pytest.mark.asyncio
async def test_user_a_cannot_update_user_b_expense(client: AsyncClient):
    a = await make_user(client, "exp_idor_upd_a@example.com")
    b = await make_user(client, "exp_idor_upd_b@example.com")
    create = await client.post("/api/v1/expenses", json=_exp(100, "food", "2026-05-02"), headers=b)
    b_expense_id = create.json()["id"]

    resp = await client.put(
        f"/api/v1/expenses/{b_expense_id}",
        json={"amount": 1},
        headers=a,
    )
    assert resp.status_code == 404


@pytest.mark.asyncio
async def test_user_a_cannot_delete_user_b_expense(client: AsyncClient):
    a = await make_user(client, "exp_idor_del_a@example.com")
    b = await make_user(client, "exp_idor_del_b@example.com")
    create = await client.post("/api/v1/expenses", json=_exp(100, "food", "2026-05-02"), headers=b)
    b_expense_id = create.json()["id"]

    resp = await client.delete(f"/api/v1/expenses/{b_expense_id}", headers=a)
    assert resp.status_code == 404


# ─── Auth ─────────────────────────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_unauth_returns_401(client: AsyncClient):
    for method, path in [
        ("GET", "/api/v1/expenses"),
        ("POST", "/api/v1/expenses"),
        ("GET", "/api/v1/expenses/summary?year=2026&month=5"),
    ]:
        resp = await client.request(method, path)
        assert resp.status_code == 401, f"{method} {path}"


# ─── UPDATE edge cases + 404 ────────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_update_all_fields_including_category_and_date(client: AsyncClient):
    headers = await make_user(client, "exp_upd_all@example.com")
    create = await client.post("/api/v1/expenses", json=_exp(50, "food", "2026-05-02", "ilk"), headers=headers)
    eid = create.json()["id"]
    resp = await client.put(
        f"/api/v1/expenses/{eid}",
        json={"category": "transport", "date": "2026-06-01", "description": "  yeni  "},
        headers=headers,
    )
    assert resp.status_code == 200
    data = resp.json()
    assert data["category"] == "transport"
    assert data["date"] == "2026-06-01"
    assert data["description"] == "yeni"


@pytest.mark.asyncio
async def test_update_description_to_empty_becomes_none(client: AsyncClient):
    headers = await make_user(client, "exp_upd_empty_desc@example.com")
    create = await client.post("/api/v1/expenses", json=_exp(50, "food", "2026-05-02", "ilk"), headers=headers)
    eid = create.json()["id"]
    resp = await client.put(f"/api/v1/expenses/{eid}", json={"description": "   "}, headers=headers)
    assert resp.status_code == 200
    assert resp.json()["description"] is None


@pytest.mark.asyncio
async def test_update_credit_card_link_and_unlink(client: AsyncClient):
    """credit_card_id explicit None ile baglanti kaldirilir (model_fields_set)."""
    headers = await make_user(client, "exp_cc_link@example.com")
    create = await client.post(
        "/api/v1/expenses",
        json={"amount": 100, "category": "food", "date": "2026-05-02", "credit_card_id": None, "is_paid": True},
        headers=headers,
    )
    eid = create.json()["id"]
    assert create.json()["is_paid"] is True
    # is_paid degistir
    resp = await client.put(f"/api/v1/expenses/{eid}", json={"is_paid": False, "credit_card_id": None}, headers=headers)
    assert resp.status_code == 200
    assert resp.json()["is_paid"] is False
    assert resp.json()["credit_card_id"] is None


@pytest.mark.asyncio
async def test_delete_nonexistent_returns_404(client: AsyncClient):
    headers = await make_user(client, "exp_del_404@example.com")
    resp = await client.delete("/api/v1/expenses/99999", headers=headers)
    assert resp.status_code == 404


# ─── Çift sayım kuralı (summary) ────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_summary_excludes_paid_credit_card_expense(client: AsyncClient):
    """credit_card_id NOT NULL + is_paid=true → ham toplamdan haric."""
    headers = await make_user(client, "exp_double@example.com")
    # Gerçek kart oluştur (credit_card_id FK; olmayan id FK ihlali → harcama oluşmaz)
    card = await client.post("/api/v1/credit-cards", json={"name": "Test Kart"}, headers=headers)
    cid = card.json()["id"]
    # Normal nakit harcama — dahil
    await client.post("/api/v1/expenses", json=_exp(500, "food", "2026-05-05"), headers=headers)
    # Kart + odendi → HARIC
    await client.post(
        "/api/v1/expenses",
        json={"amount": 1000, "category": "food", "date": "2026-05-06", "credit_card_id": cid, "is_paid": True},
        headers=headers,
    )
    # Kart ama odenmemis → dahil
    await client.post(
        "/api/v1/expenses",
        json={"amount": 200, "category": "food", "date": "2026-05-07", "credit_card_id": cid, "is_paid": False},
        headers=headers,
    )
    resp = await client.get("/api/v1/expenses/summary?year=2026&month=5", headers=headers)
    data = resp.json()
    # 500 + 200 = 700 (1000 haric)
    assert float(data["total"]) == 700.0
    assert data["count"] == 2


# ─── EXPORT / IMPORT ────────────────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_export_xlsx(client: AsyncClient):
    headers = await make_user(client, "exp_export@example.com")
    await client.post("/api/v1/expenses", json=_exp(150, "groceries", "2026-05-02", "market"), headers=headers)
    resp = await client.get("/api/v1/expenses/export", headers=headers)
    assert resp.status_code == 200
    assert "spreadsheetml" in resp.headers["content-type"]
    assert "harcamalar.xlsx" in resp.headers["content-disposition"]
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    assert ws.cell(row=1, column=1).value == "Tarih"
    assert ws.cell(row=2, column=2).value == "groceries"


@pytest.mark.asyncio
async def test_export_filtered_by_month(client: AsyncClient):
    headers = await make_user(client, "exp_export_filter@example.com")
    await client.post("/api/v1/expenses", json=_exp(100, "food", "2026-05-01"), headers=headers)
    await client.post("/api/v1/expenses", json=_exp(200, "food", "2026-04-01"), headers=headers)
    resp = await client.get("/api/v1/expenses/export?year=2026&month=5", headers=headers)
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    assert wb.active.max_row == 2  # header + 1


@pytest.mark.asyncio
async def test_import_xlsx(client: AsyncClient):
    headers = await make_user(client, "exp_import@example.com")
    content = _excel_bytes(
        [
            ("2026-05-01", "Market", 150.5, "Migros"),
            ("2026-05-02", "ulaşım", 50, None),
            ("2026-05-03", "food", "75", "yemek"),
        ]
    )
    resp = await client.post(
        "/api/v1/expenses/import",
        files={"file": ("harcama.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=headers,
    )
    assert resp.status_code == 201
    data = resp.json()
    assert len(data) == 3
    cats = sorted(d["category"] for d in data)
    assert cats == ["food", "groceries", "transport"]


@pytest.mark.asyncio
async def test_import_skips_invalid_rows(client: AsyncClient):
    headers = await make_user(client, "exp_import_skip@example.com")
    content = _excel_bytes(
        [
            ("2026-05-01", "Market", 150, "ok"),
            (None, None, None, None),
            ("bad-date", "Market", 100, "kotu tarih"),
            ("2026-05-02", "yok-kategori", 100, "kotu kategori"),
            ("2026-05-03", "Market", -5, "negatif"),
            ("2026-05-04", "Market", "xyz", "tutar hatasi"),
        ]
    )
    resp = await client.post(
        "/api/v1/expenses/import",
        files={"file": ("harcama.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=headers,
    )
    assert resp.status_code == 201
    assert len(resp.json()) == 1


@pytest.mark.asyncio
async def test_import_rejects_non_excel(client: AsyncClient):
    headers = await make_user(client, "exp_import_bad@example.com")
    resp = await client.post(
        "/api/v1/expenses/import",
        files={"file": ("harcama.txt", b"hello", "text/plain")},
        headers=headers,
    )
    assert resp.status_code == 422


@pytest.mark.asyncio
async def test_import_rejects_fake_magic(client: AsyncClient):
    headers = await make_user(client, "exp_import_magic@example.com")
    resp = await client.post(
        "/api/v1/expenses/import",
        files={"file": ("harcama.xlsx", b"NOT-A-ZIP", "application/octet-stream")},
        headers=headers,
    )
    assert resp.status_code == 422
