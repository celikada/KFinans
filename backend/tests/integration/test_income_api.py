"""Income CRUD + summary endpoint testleri."""

import io
from datetime import date as _date
from datetime import datetime as _datetime
from zoneinfo import ZoneInfo

import openpyxl
import pytest
from httpx import AsyncClient

from tests.conftest import make_user

# Endpoint realize tarihini Europe/Istanbul ile hesaplar (income.py); test de aynı
# tz'yi kullanmalı, yoksa UTC↔Istanbul gece yarısı farkında off-by-one olur.
_ISTANBUL = ZoneInfo("Europe/Istanbul")


def _monthly_past_count(start_iso: str, day_of_month: int = 1) -> int:
    """start_iso'dan bugüne aylık geçmiş dönem sayısı (tarih + tz bağımsız).

    realize-past, period_date <= today (Istanbul) olan aylık dönemleri üretir.
    day_of_month <= bugünün günü ise içinde bulunulan ayın dönemi de geçmiştir.
    Testler sabit gün sayısı yazamaz (ay dönümünde kayar) — dinamik hesapla.
    """
    start = _date.fromisoformat(start_iso)
    today = _datetime.now(_ISTANBUL).date()
    count = (today.year - start.year) * 12 + (today.month - start.month)
    if day_of_month <= today.day:
        count += 1
    return max(count, 0)


def _inc(amount: float, category: str, date: str, description: str | None = None) -> dict:
    return {"amount": amount, "category": category, "date": date, "description": description}


def _recurring(
    title: str = "Maaş",
    amount: float = 50000.0,
    category: str = "salary",
    recurrence: str = "monthly",
    start_date: str = "2026-01-01",
    day_of_month: int = 1,
    months: list[int] | None = None,
    end_date: str | None = None,
) -> dict:
    payload: dict = {
        "title": title,
        "amount": amount,
        "category": category,
        "recurrence": recurrence,
        "start_date": start_date,
        "day_of_month": day_of_month,
    }
    if months is not None:
        payload["months"] = months
    if end_date is not None:
        payload["end_date"] = end_date
    return payload


def _excel_bytes(rows: list[tuple]) -> bytes:
    """Header + verilen satirlardan .xlsx bytes uretir (import testleri icin)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(("Tarih", "Kategori", "Tutar", "Açıklama"))
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.mark.asyncio
async def test_empty_list(client: AsyncClient):
    headers = await make_user(client, "inc_empty@example.com")
    resp = await client.get("/api/v1/income", headers=headers)
    assert resp.status_code == 200
    assert resp.json() == []


@pytest.mark.asyncio
async def test_create_income(client: AsyncClient):
    headers = await make_user(client, "inc_create@example.com")
    resp = await client.post(
        "/api/v1/income",
        json=_inc(50000, "salary", "2026-05-01", "Mayıs maaşı"),
        headers=headers,
    )
    assert resp.status_code == 201
    data = resp.json()
    assert float(data["amount"]) == 50000.0
    assert data["category"] == "salary"
    assert data["date"] == "2026-05-01"
    assert data["description"] == "Mayıs maaşı"
    assert "id" in data


@pytest.mark.asyncio
async def test_create_invalid_category(client: AsyncClient):
    headers = await make_user(client, "inc_badcat@example.com")
    resp = await client.post("/api/v1/income", json=_inc(1000, "crypto", "2026-05-01"), headers=headers)
    assert resp.status_code == 422


@pytest.mark.asyncio
async def test_create_negative_amount(client: AsyncClient):
    headers = await make_user(client, "inc_neg@example.com")
    resp = await client.post("/api/v1/income", json=_inc(-500, "salary", "2026-05-01"), headers=headers)
    assert resp.status_code == 422


@pytest.mark.asyncio
async def test_list_filter_by_month(client: AsyncClient):
    headers = await make_user(client, "inc_filter@example.com")
    await client.post("/api/v1/income", json=_inc(50000, "salary", "2026-05-01"), headers=headers)
    await client.post("/api/v1/income", json=_inc(10000, "bonus", "2026-04-15"), headers=headers)

    resp = await client.get("/api/v1/income?year=2026&month=5", headers=headers)
    data = resp.json()
    assert len(data) == 1
    assert data[0]["category"] == "salary"


@pytest.mark.asyncio
async def test_update_income(client: AsyncClient):
    headers = await make_user(client, "inc_update@example.com")
    create = await client.post("/api/v1/income", json=_inc(50000, "salary", "2026-05-01"), headers=headers)
    inc_id = create.json()["id"]

    resp = await client.put(f"/api/v1/income/{inc_id}", json={"amount": 55000.0}, headers=headers)
    assert resp.status_code == 200
    assert float(resp.json()["amount"]) == 55000.0


@pytest.mark.asyncio
async def test_delete_income(client: AsyncClient):
    headers = await make_user(client, "inc_delete@example.com")
    create = await client.post("/api/v1/income", json=_inc(50000, "salary", "2026-05-01"), headers=headers)
    inc_id = create.json()["id"]

    resp = await client.delete(f"/api/v1/income/{inc_id}", headers=headers)
    assert resp.status_code == 204

    list_resp = await client.get("/api/v1/income", headers=headers)
    assert list_resp.json() == []


@pytest.mark.asyncio
async def test_idor_update(client: AsyncClient):
    h1 = await make_user(client, "inc_idor1@example.com")
    h2 = await make_user(client, "inc_idor2@example.com")
    create = await client.post("/api/v1/income", json=_inc(50000, "salary", "2026-05-01"), headers=h1)
    inc_id = create.json()["id"]

    resp = await client.put(f"/api/v1/income/{inc_id}", json={"amount": 1.0}, headers=h2)
    assert resp.status_code == 404


@pytest.mark.asyncio
async def test_idor_delete(client: AsyncClient):
    h1 = await make_user(client, "inc_idor3@example.com")
    h2 = await make_user(client, "inc_idor4@example.com")
    create = await client.post("/api/v1/income", json=_inc(50000, "salary", "2026-05-01"), headers=h1)
    inc_id = create.json()["id"]

    resp = await client.delete(f"/api/v1/income/{inc_id}", headers=h2)
    assert resp.status_code == 404


@pytest.mark.asyncio
async def test_unauthenticated(client: AsyncClient):
    resp = await client.get("/api/v1/income")
    assert resp.status_code == 401


@pytest.mark.asyncio
async def test_summary_empty(client: AsyncClient):
    headers = await make_user(client, "inc_sum_empty@example.com")
    resp = await client.get("/api/v1/income/summary?year=2026&month=5", headers=headers)
    assert resp.status_code == 200
    data = resp.json()
    assert float(data["total"]) == 0.0
    assert data["count"] == 0
    assert data["by_category"] == []


@pytest.mark.asyncio
async def test_summary_with_data(client: AsyncClient):
    headers = await make_user(client, "inc_sum_data@example.com")
    await client.post("/api/v1/income", json=_inc(50000, "salary", "2026-05-01"), headers=headers)
    await client.post("/api/v1/income", json=_inc(10000, "freelance", "2026-05-15"), headers=headers)
    await client.post("/api/v1/income", json=_inc(5000, "dividend", "2026-05-20"), headers=headers)
    # Farklı ay — summary'e dahil olmamalı
    await client.post("/api/v1/income", json=_inc(20000, "bonus", "2026-04-01"), headers=headers)

    resp = await client.get("/api/v1/income/summary?year=2026&month=5", headers=headers)
    data = resp.json()
    assert float(data["total"]) == 65000.0
    assert data["count"] == 3
    categories = [b["category"] for b in data["by_category"]]
    assert "salary" in categories
    assert "freelance" in categories


@pytest.mark.asyncio
async def test_net_balance_calculation(client: AsyncClient):
    """Gelir ve gider endpoint'leri birlikte çalışıyor mu — net bakiye kontrolü."""
    headers = await make_user(client, "inc_net@example.com")
    await client.post("/api/v1/income", json=_inc(50000, "salary", "2026-05-01"), headers=headers)
    await client.post(
        "/api/v1/expenses",
        json={"amount": 20000, "category": "bills", "date": "2026-05-10"},
        headers=headers,
    )

    income_sum = await client.get("/api/v1/income/summary?year=2026&month=5", headers=headers)
    expense_sum = await client.get("/api/v1/expenses/summary?year=2026&month=5", headers=headers)

    net = float(income_sum.json()["total"]) - float(expense_sum.json()["total"])
    assert net == 30000.0


# ─── UPDATE/DELETE edge cases ──────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_update_all_fields(client: AsyncClient):
    headers = await make_user(client, "inc_upd_all@example.com")
    create = await client.post("/api/v1/income", json=_inc(1000, "salary", "2026-05-01", "ilk"), headers=headers)
    inc_id = create.json()["id"]
    resp = await client.put(
        f"/api/v1/income/{inc_id}",
        json={"amount": 2000.0, "category": "bonus", "date": "2026-06-15", "description": "  yeni  "},
        headers=headers,
    )
    assert resp.status_code == 200
    data = resp.json()
    assert float(data["amount"]) == 2000.0
    assert data["category"] == "bonus"
    assert data["date"] == "2026-06-15"
    assert data["description"] == "yeni"


@pytest.mark.asyncio
async def test_update_nonexistent_returns_404(client: AsyncClient):
    headers = await make_user(client, "inc_upd_404@example.com")
    resp = await client.put("/api/v1/income/99999", json={"amount": 100.0}, headers=headers)
    assert resp.status_code == 404


@pytest.mark.asyncio
async def test_delete_nonexistent_returns_404(client: AsyncClient):
    headers = await make_user(client, "inc_del_404@example.com")
    resp = await client.delete("/api/v1/income/99999", headers=headers)
    assert resp.status_code == 404


@pytest.mark.asyncio
async def test_list_invalid_category_returns_422(client: AsyncClient):
    headers = await make_user(client, "inc_list_badcat@example.com")
    resp = await client.get("/api/v1/income?category=invalid", headers=headers)
    assert resp.status_code == 422


@pytest.mark.asyncio
async def test_list_filter_by_category(client: AsyncClient):
    headers = await make_user(client, "inc_list_cat@example.com")
    await client.post("/api/v1/income", json=_inc(1000, "salary", "2026-05-01"), headers=headers)
    await client.post("/api/v1/income", json=_inc(2000, "rental", "2026-05-02"), headers=headers)
    resp = await client.get("/api/v1/income?category=rental", headers=headers)
    data = resp.json()
    assert len(data) == 1
    assert data[0]["category"] == "rental"


# ─── EXPORT / IMPORT ────────────────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_export_xlsx(client: AsyncClient):
    headers = await make_user(client, "inc_export@example.com")
    await client.post("/api/v1/income", json=_inc(50000, "salary", "2026-05-01", "maaş"), headers=headers)
    resp = await client.get("/api/v1/income/export", headers=headers)
    assert resp.status_code == 200
    assert "spreadsheetml" in resp.headers["content-type"]
    assert "gelirler.xlsx" in resp.headers["content-disposition"]
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    assert ws.cell(row=1, column=1).value == "Tarih"
    assert ws.cell(row=2, column=2).value == "salary"


@pytest.mark.asyncio
async def test_export_filtered_by_month(client: AsyncClient):
    headers = await make_user(client, "inc_export_filter@example.com")
    await client.post("/api/v1/income", json=_inc(100, "salary", "2026-05-01"), headers=headers)
    await client.post("/api/v1/income", json=_inc(200, "salary", "2026-04-01"), headers=headers)
    resp = await client.get("/api/v1/income/export?year=2026&month=5", headers=headers)
    assert resp.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    # 1 header + 1 veri satiri
    assert ws.max_row == 2


@pytest.mark.asyncio
async def test_import_xlsx(client: AsyncClient):
    headers = await make_user(client, "inc_import@example.com")
    content = _excel_bytes(
        [
            ("2026-05-01", "Maaş", 50000, "Mayıs"),
            ("2026-05-15", "freelance", 7500.5, "Proje"),
            ("2026-05-20", "Temettü", "1200", None),
        ]
    )
    resp = await client.post(
        "/api/v1/income/import",
        files={"file": ("gelir.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=headers,
    )
    assert resp.status_code == 201
    data = resp.json()
    assert len(data) == 3
    cats = sorted(d["category"] for d in data)
    assert cats == ["dividend", "freelance", "salary"]


@pytest.mark.asyncio
async def test_import_skips_invalid_rows(client: AsyncClient):
    headers = await make_user(client, "inc_import_skip@example.com")
    content = _excel_bytes(
        [
            ("2026-05-01", "Maaş", 50000, "ok"),
            (None, None, None, None),  # bos satir
            ("not-a-date", "Maaş", 100, "kotu tarih"),
            ("2026-05-02", "bilinmeyen", 100, "kotu kategori"),
            ("2026-05-03", "Maaş", -50, "negatif"),
            ("2026-05-04", "Maaş", "abc", "tutar parse hatasi"),
        ]
    )
    resp = await client.post(
        "/api/v1/income/import",
        files={"file": ("gelir.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=headers,
    )
    assert resp.status_code == 201
    # Sadece ilk gecerli satir
    assert len(resp.json()) == 1


@pytest.mark.asyncio
async def test_import_rejects_non_excel(client: AsyncClient):
    headers = await make_user(client, "inc_import_bad@example.com")
    resp = await client.post(
        "/api/v1/income/import",
        files={"file": ("gelir.txt", b"hello", "text/plain")},
        headers=headers,
    )
    assert resp.status_code == 422


@pytest.mark.asyncio
async def test_import_rejects_fake_magic(client: AsyncClient):
    """xlsx uzantili ama PDF icerikli — magic byte uyusmaz."""
    headers = await make_user(client, "inc_import_magic@example.com")
    resp = await client.post(
        "/api/v1/income/import",
        files={"file": ("gelir.xlsx", b"%PDF-1.7 fake", "application/octet-stream")},
        headers=headers,
    )
    assert resp.status_code == 422


# ─── RECURRING CRUD ─────────────────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_recurring_empty_list(client: AsyncClient):
    headers = await make_user(client, "rec_empty@example.com")
    resp = await client.get("/api/v1/income/recurring", headers=headers)
    assert resp.status_code == 200
    assert resp.json() == []


@pytest.mark.asyncio
async def test_recurring_create_monthly(client: AsyncClient):
    headers = await make_user(client, "rec_create@example.com")
    resp = await client.post("/api/v1/income/recurring", json=_recurring(), headers=headers)
    assert resp.status_code == 201
    data = resp.json()
    assert data["title"] == "Maaş"
    assert data["recurrence"] == "monthly"
    assert float(data["amount"]) == 50000.0


@pytest.mark.asyncio
async def test_recurring_create_custom_requires_months(client: AsyncClient):
    headers = await make_user(client, "rec_custom_nomonths@example.com")
    payload = _recurring(recurrence="custom")  # months yok
    resp = await client.post("/api/v1/income/recurring", json=payload, headers=headers)
    assert resp.status_code == 422


@pytest.mark.asyncio
async def test_recurring_create_custom_with_months(client: AsyncClient):
    headers = await make_user(client, "rec_custom_ok@example.com")
    payload = _recurring(recurrence="custom", months=[3, 6, 9])
    resp = await client.post("/api/v1/income/recurring", json=payload, headers=headers)
    assert resp.status_code == 201
    assert resp.json()["months"] == [3, 6, 9]


@pytest.mark.asyncio
async def test_recurring_update(client: AsyncClient):
    headers = await make_user(client, "rec_update@example.com")
    create = await client.post("/api/v1/income/recurring", json=_recurring(), headers=headers)
    rid = create.json()["id"]
    resp = await client.put(
        f"/api/v1/income/recurring/{rid}",
        json={"amount": 60000.0, "notes": "zam"},
        headers=headers,
    )
    assert resp.status_code == 200
    data = resp.json()
    assert float(data["amount"]) == 60000.0
    assert data["notes"] == "zam"
    assert data["title"] == "Maaş"  # degismedi


@pytest.mark.asyncio
async def test_recurring_update_404(client: AsyncClient):
    headers = await make_user(client, "rec_upd_404@example.com")
    resp = await client.put("/api/v1/income/recurring/99999", json={"amount": 1.0}, headers=headers)
    assert resp.status_code == 404


@pytest.mark.asyncio
async def test_recurring_delete(client: AsyncClient):
    headers = await make_user(client, "rec_delete@example.com")
    create = await client.post("/api/v1/income/recurring", json=_recurring(), headers=headers)
    rid = create.json()["id"]
    resp = await client.delete(f"/api/v1/income/recurring/{rid}", headers=headers)
    assert resp.status_code == 204
    list_resp = await client.get("/api/v1/income/recurring", headers=headers)
    assert list_resp.json() == []


@pytest.mark.asyncio
async def test_recurring_delete_404(client: AsyncClient):
    headers = await make_user(client, "rec_del_404@example.com")
    resp = await client.delete("/api/v1/income/recurring/99999", headers=headers)
    assert resp.status_code == 404


@pytest.mark.asyncio
async def test_recurring_idor_update(client: AsyncClient):
    h1 = await make_user(client, "rec_idor1@example.com")
    h2 = await make_user(client, "rec_idor2@example.com")
    create = await client.post("/api/v1/income/recurring", json=_recurring(), headers=h1)
    rid = create.json()["id"]
    resp = await client.put(f"/api/v1/income/recurring/{rid}", json={"amount": 1.0}, headers=h2)
    assert resp.status_code == 404
    resp2 = await client.delete(f"/api/v1/income/recurring/{rid}", headers=h2)
    assert resp2.status_code == 404


# ─── DASHBOARD ──────────────────────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_dashboard_empty(client: AsyncClient):
    headers = await make_user(client, "dash_empty@example.com")
    resp = await client.get("/api/v1/income/dashboard?year=2026&month=5", headers=headers)
    assert resp.status_code == 200
    data = resp.json()
    assert float(data["this_month_actual"]) == 0.0
    assert float(data["ytd_actual"]) == 0.0
    assert float(data["this_month_recurring"]) == 0.0
    assert float(data["year_total_estimate"]) == 0.0


@pytest.mark.asyncio
async def test_dashboard_actual_and_recurring(client: AsyncClient):
    headers = await make_user(client, "dash_full@example.com")
    # Gerceklesen: Ocak + Mayis
    await client.post("/api/v1/income", json=_inc(10000, "salary", "2026-01-15"), headers=headers)
    await client.post("/api/v1/income", json=_inc(20000, "salary", "2026-05-10"), headers=headers)
    # Aylik recurring 12 ay (Ocak basindan)
    await client.post(
        "/api/v1/income/recurring",
        json=_recurring(amount=5000.0, recurrence="monthly", start_date="2026-01-01"),
        headers=headers,
    )
    resp = await client.get("/api/v1/income/dashboard?year=2026&month=5", headers=headers)
    data = resp.json()
    assert float(data["this_month_actual"]) == 20000.0
    assert float(data["ytd_actual"]) == 30000.0
    assert float(data["this_month_recurring"]) == 5000.0
    # Ocak..Mayis = 5 ay aktif
    assert float(data["ytd_recurring"]) == 25000.0
    # Haziran..Aralik = 7 ay
    assert float(data["remaining_year_recurring"]) == 35000.0
    # ytd_actual + remaining
    assert float(data["year_total_estimate"]) == 65000.0


@pytest.mark.asyncio
async def test_dashboard_recurrence_variants(client: AsyncClient):
    """quarterly/biannual/yearly/custom/one_time _applies_in_month dallari."""
    headers = await make_user(client, "dash_variants@example.com")
    await client.post(
        "/api/v1/income/recurring",
        json=_recurring(title="3aylik", amount=300, recurrence="quarterly", start_date="2026-01-01"),
        headers=headers,
    )
    await client.post(
        "/api/v1/income/recurring",
        json=_recurring(title="6aylik", amount=600, recurrence="biannual", start_date="2026-01-01"),
        headers=headers,
    )
    await client.post(
        "/api/v1/income/recurring",
        json=_recurring(title="yillik", amount=1200, recurrence="yearly", start_date="2026-03-01"),
        headers=headers,
    )
    await client.post(
        "/api/v1/income/recurring",
        json=_recurring(title="ozel", amount=900, recurrence="custom", months=[2, 8], start_date="2026-01-01"),
        headers=headers,
    )
    await client.post(
        "/api/v1/income/recurring",
        json=_recurring(title="tekseferlik", amount=5000, recurrence="one_time", start_date="2026-07-01"),
        headers=headers,
    )
    resp = await client.get("/api/v1/income/dashboard?year=2026&month=12", headers=headers)
    data = resp.json()
    # quarterly: Oca,Nis,Tem,Eki = 4*300=1200; biannual: Oca,Tem=2*600=1200;
    # yearly Mart=1200; custom Sub,Agu=2*900=1800; one_time Tem=5000
    assert float(data["ytd_recurring"]) == 1200 + 1200 + 1200 + 1800 + 5000


@pytest.mark.asyncio
async def test_dashboard_unauthenticated(client: AsyncClient):
    resp = await client.get("/api/v1/income/dashboard?year=2026&month=5")
    assert resp.status_code == 401


# ─── REALIZE ────────────────────────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_realize_single_period(client: AsyncClient):
    headers = await make_user(client, "realize_single@example.com")
    create = await client.post(
        "/api/v1/income/recurring",
        json=_recurring(amount=50000, recurrence="monthly", start_date="2026-01-01", day_of_month=1),
        headers=headers,
    )
    rid = create.json()["id"]
    resp = await client.post(
        f"/api/v1/income/recurring/{rid}/realize",
        json={"year": 2026, "month": 2},
        headers=headers,
    )
    assert resp.status_code == 200
    data = resp.json()
    assert data["realized"] == 1
    assert data["skipped"] == 0
    assert len(data["income_ids"]) == 1
    # income listesinde gorunmeli
    listed = await client.get("/api/v1/income?year=2026&month=2", headers=headers)
    body = listed.json()
    assert len(body) == 1
    assert body[0]["recurring_income_id"] == rid


@pytest.mark.asyncio
async def test_realize_idempotent_skip(client: AsyncClient):
    headers = await make_user(client, "realize_idem@example.com")
    create = await client.post(
        "/api/v1/income/recurring",
        json=_recurring(recurrence="monthly", start_date="2026-01-01", day_of_month=1),
        headers=headers,
    )
    rid = create.json()["id"]
    first = await client.post(f"/api/v1/income/recurring/{rid}/realize", json={"year": 2026, "month": 2}, headers=headers)
    assert first.json()["realized"] == 1
    second = await client.post(f"/api/v1/income/recurring/{rid}/realize", json={"year": 2026, "month": 2}, headers=headers)
    assert second.json()["realized"] == 0
    assert second.json()["skipped"] == 1


@pytest.mark.asyncio
async def test_realize_404(client: AsyncClient):
    headers = await make_user(client, "realize_404@example.com")
    resp = await client.post("/api/v1/income/recurring/99999/realize", json={"year": 2026, "month": 2}, headers=headers)
    assert resp.status_code == 404


@pytest.mark.asyncio
async def test_realize_out_of_period_422(client: AsyncClient):
    """Periyot disinda bir ay icin realize → 422."""
    headers = await make_user(client, "realize_oop@example.com")
    create = await client.post(
        "/api/v1/income/recurring",
        json=_recurring(recurrence="yearly", start_date="2026-03-01", day_of_month=1),
        headers=headers,
    )
    rid = create.json()["id"]
    # yearly Mart kaydi icin Subat istegi periyot disi
    resp = await client.post(f"/api/v1/income/recurring/{rid}/realize", json={"year": 2026, "month": 2}, headers=headers)
    assert resp.status_code == 422


@pytest.mark.asyncio
async def test_realize_future_date_422(client: AsyncClient):
    """Odeme gunu henuz gelmedi (bugun 2026-05-31) → 422."""
    headers = await make_user(client, "realize_future@example.com")
    create = await client.post(
        "/api/v1/income/recurring",
        json=_recurring(recurrence="monthly", start_date="2026-01-01", day_of_month=1),
        headers=headers,
    )
    rid = create.json()["id"]
    resp = await client.post(f"/api/v1/income/recurring/{rid}/realize", json={"year": 2030, "month": 1}, headers=headers)
    assert resp.status_code == 422


@pytest.mark.asyncio
async def test_realize_past_all_periods(client: AsyncClient):
    """realize-past start_date'ten bugune tum gecmis donemleri olusturur."""
    headers = await make_user(client, "realize_past@example.com")
    create = await client.post(
        "/api/v1/income/recurring",
        json=_recurring(recurrence="monthly", start_date="2026-01-01", day_of_month=1),
        headers=headers,
    )
    rid = create.json()["id"]
    resp = await client.post(f"/api/v1/income/recurring/{rid}/realize-past", headers=headers)
    assert resp.status_code == 200
    data = resp.json()
    # 2026-01-01'den bugüne aylık dönemler (tarih-bağımsız hesap; ay dönümünde kaymaz)
    expected = _monthly_past_count("2026-01-01", day_of_month=1)
    assert data["realized"] == expected
    # Tekrar cagir → hepsi skip
    again = await client.post(f"/api/v1/income/recurring/{rid}/realize-past", headers=headers)
    assert again.json()["realized"] == 0
    assert again.json()["skipped"] == expected


@pytest.mark.asyncio
async def test_realize_past_404(client: AsyncClient):
    headers = await make_user(client, "realize_past_404@example.com")
    resp = await client.post("/api/v1/income/recurring/99999/realize-past", headers=headers)
    assert resp.status_code == 404


@pytest.mark.asyncio
async def test_realize_all_past(client: AsyncClient):
    """realize-all-past tum recurring'ler icin gecmis donemleri olusturur."""
    headers = await make_user(client, "realize_all@example.com")
    await client.post(
        "/api/v1/income/recurring",
        json=_recurring(title="maas", recurrence="monthly", start_date="2026-04-01", day_of_month=1),
        headers=headers,
    )
    await client.post(
        "/api/v1/income/recurring",
        json=_recurring(title="kira", category="rental", recurrence="monthly", start_date="2026-05-01", day_of_month=1),
        headers=headers,
    )
    resp = await client.post("/api/v1/income/recurring/realize-all-past", headers=headers)
    assert resp.status_code == 200
    # maas (2026-04-01) + kira (2026-05-01) aylık dönemler (tarih-bağımsız)
    expected = _monthly_past_count("2026-04-01") + _monthly_past_count("2026-05-01")
    assert resp.json()["realized"] == expected


@pytest.mark.asyncio
async def test_realize_unauthenticated(client: AsyncClient):
    resp = await client.post("/api/v1/income/recurring/realize-all-past")
    assert resp.status_code == 401


# ─── Periyodik gelir dönem yönetimi + realize/skip geri alma ────────────────


async def _create_recurring(client: AsyncClient, headers: dict, **kw) -> int:
    resp = await client.post("/api/v1/income/recurring", json=_recurring(**kw), headers=headers)
    assert resp.status_code == 201
    return resp.json()["id"]


@pytest.mark.asyncio
async def test_recurring_periods_lists_status(client: AsyncClient):
    """GET /income/recurring/{id}/periods her dönemi pending/realized/skipped döner."""
    headers = await make_user(client, "inc_periods@example.com")
    rid = await _create_recurring(client, headers, start_date="2026-01-01", day_of_month=1)

    r = await client.get(f"/api/v1/income/recurring/{rid}/periods", headers=headers)
    assert r.status_code == 200
    periods = r.json()["periods"]
    assert len(periods) >= 2
    assert all(p["status"] == "pending" for p in periods)

    rr = await client.post(
        f"/api/v1/income/recurring/{rid}/realize",
        json={"year": 2026, "month": 2},
        headers=headers,
    )
    assert rr.status_code == 200 and rr.json()["realized"] == 1

    sk = await client.post(
        "/api/v1/recurring/skips",
        json={"kind": "income", "ref_id": rid, "year": 2026, "month": 3},
        headers=headers,
    )
    assert sk.status_code == 201

    r2 = await client.get(f"/api/v1/income/recurring/{rid}/periods", headers=headers)
    by_month = {(p["year"], p["month"]): p for p in r2.json()["periods"]}
    assert by_month[(2026, 2)]["status"] == "realized"
    assert by_month[(2026, 2)]["income_id"] is not None
    assert by_month[(2026, 3)]["status"] == "skipped"
    assert by_month[(2026, 3)]["skip_id"] is not None


@pytest.mark.asyncio
async def test_recurring_unrealize_removes_income(client: AsyncClient):
    """POST /income/recurring/{id}/unrealize gelir realize'ini geri alir (idempotent)."""
    headers = await make_user(client, "inc_unrealize@example.com")
    rid = await _create_recurring(client, headers, start_date="2026-01-01", day_of_month=1)

    await client.post(
        f"/api/v1/income/recurring/{rid}/realize",
        json={"year": 2026, "month": 2},
        headers=headers,
    )
    lst = await client.get("/api/v1/income?year=2026&month=2", headers=headers)
    assert len(lst.json()) >= 1

    un = await client.post(
        f"/api/v1/income/recurring/{rid}/unrealize",
        json={"year": 2026, "month": 2},
        headers=headers,
    )
    assert un.status_code == 200 and un.json()["removed"] == 1

    r = await client.get(f"/api/v1/income/recurring/{rid}/periods", headers=headers)
    by_month = {(p["year"], p["month"]): p for p in r.json()["periods"]}
    assert by_month[(2026, 2)]["status"] == "pending"

    un2 = await client.post(
        f"/api/v1/income/recurring/{rid}/unrealize",
        json={"year": 2026, "month": 2},
        headers=headers,
    )
    assert un2.json()["removed"] == 0


@pytest.mark.asyncio
async def test_recurring_periods_idor(client: AsyncClient):
    headers_a = await make_user(client, "inc_periods_a@example.com")
    headers_b = await make_user(client, "inc_periods_b@example.com")
    rid = await _create_recurring(client, headers_a, start_date="2026-01-01")
    r = await client.get(f"/api/v1/income/recurring/{rid}/periods", headers=headers_b)
    assert r.status_code == 404


@pytest.mark.asyncio
async def test_recurring_unrealize_unauthenticated(client: AsyncClient):
    resp = await client.post("/api/v1/income/recurring/1/unrealize", json={"year": 2026, "month": 2})
    assert resp.status_code == 401
